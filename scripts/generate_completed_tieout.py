"""Generate completed tie-out worksheet Excel workbook for the approved Sample-Data Fallback Pilot."""

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def create_completed_tieout():
    wb = openpyxl.Workbook()
    ws_summary = wb.active
    ws_summary.title = "Tie-Out Summary"
    ws_diff = wb.create_sheet(title="Difference Classification Log")
    ws_signoff = wb.create_sheet(title="Sign-Off & Approvals")

    font_title = Font(name="Calibri", size=15, bold=True, color="1F497D")
    font_sub = Font(name="Calibri", size=10, italic=True, color="595959")
    font_notice = Font(name="Calibri", size=9, bold=True, color="C00000")
    font_sec = Font(name="Calibri", size=11, bold=True, color="1F497D")
    font_tbl_hdr = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_data = Font(name="Calibri", size=10)
    font_bold = Font(name="Calibri", size=10, bold=True)
    font_blocked = Font(name="Calibri", size=10, bold=True, color="9C0006")

    fill_hdr = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    fill_sec = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
    fill_notice = PatternFill(start_color="FDE9D9", end_color="FDE9D9", fill_type="solid")
    fill_pass = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    fill_blocked = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    notice_text = (
        "PILOT FALLBACK LIMITATION NOTICE (RISK-002 / OQ-014 / DEC-REQ-01):\n"
        "This pilot run and tie-out worksheet use synthetic sample data (d365_gl_actuals.csv, budget_fy26.csv) "
        "rather than sanitized real client month-end data per Project Owner approval. All figures, variances, exception findings, "
        "and forecast scenarios presented herein are illustrative and intended solely for software validation, UAT familiarization, "
        "and workflow rehearsal. They do not constitute a formal production sign-off or audit conclusion until sanitized real client data is ingested."
    )

    # ----------------- SHEET 1: Summary -----------------
    ws_summary.merge_cells("A1:G1")
    ws_summary["A1"] = "FP&A Month-End Pilot Tie-Out Worksheet (GATE-13 - APPROVED FALLBACK)"
    ws_summary["A1"].font = font_title

    ws_summary.merge_cells("A2:G2")
    ws_summary["A2"] = (
        "Per Doc 28 §4.4 & §4.6 - Execution of Sample-Data Fallback Pilot against baseline expectations"
    )
    ws_summary["A2"].font = font_sub

    ws_summary.merge_cells("A3:G4")
    ws_summary["A3"] = notice_text
    ws_summary["A3"].font = font_notice
    ws_summary["A3"].fill = fill_notice
    ws_summary["A3"].alignment = Alignment(wrap_text=True, vertical="center")

    meta = [
        (
            "Client / Project Name:",
            "Northwind Industries — FY26 Fallback Pilot (Approved Sample Data)",
        ),
        ("Period Covered:", "2026-P09 (FY26-P09)"),
        (
            "Build Under Test:",
            "v0.1.0 — Inno Setup SHA-256: e871404c0003f905d4baeeae33ce7bcf14cefbdf04803db17f1bf249b6b9dcf7",
        ),
        ("Execution Date:", "2026-10-03"),
        ("Fallback Authorization:", "Approved by Project Owner (DEC-053 / RISK-002 / GATE-13)"),
    ]
    for r_idx, (k, v) in enumerate(meta, start=6):
        ws_summary.cell(r_idx, 1, k).font = font_bold
        ws_summary.cell(r_idx, 2, v).font = font_data

    # Section 1: Ingestion
    ws_summary.cell(12, 1, "1. Imported Files Summary").font = font_sec
    headers_files = [
        "File Name",
        "Source Type",
        "Batch ID",
        "Row Count",
        "SHA-256 (Truncated)",
        "Data Quality Score",
        "Ingestion Status",
    ]
    for c_idx, h in enumerate(headers_files, start=1):
        c = ws_summary.cell(13, c_idx, h)
        c.font = font_tbl_hdr
        c.fill = fill_hdr
        c.alignment = Alignment(horizontal="center")

    files_data = [
        (
            "d365_gl_actuals.csv",
            "D365 GL Actuals",
            "BATCH-001",
            250037,
            "2c791e5f27...",
            "84.0% (Failed IMP-023)",
            "REJECTED / UNCOMMITTED",
        ),
        (
            "budget_fy26.csv",
            "Budget FY26",
            "BATCH-002",
            1980,
            "74972d7f8d...",
            "100.0% (32 checks passed)",
            "COMMITTED",
        ),
    ]
    for r_idx, row in enumerate(files_data, start=14):
        for c_idx, val in enumerate(row, start=1):
            cell = ws_summary.cell(r_idx, c_idx, val)
            cell.font = font_data
            cell.border = thin_border
            if c_idx == 7:
                cell.fill = fill_blocked if row[0] == "d365_gl_actuals.csv" else fill_pass
                cell.font = font_blocked if row[0] == "d365_gl_actuals.csv" else font_data
                cell.alignment = Alignment(horizontal="center")

    # Section 2: BvA Totals
    ws_summary.cell(
        17, 1, "2. Budget vs Actuals (BvA) Financial Comparison (MTD: 2026-P09)"
    ).font = font_sec
    headers_bva = [
        "Account Code",
        "Statement Line / Category",
        "Actual Amount (₹)",
        "Budget Amount (₹)",
        "Variance Amount (₹)",
        "Variance %",
        "Status / Notes",
    ]
    for c_idx, h in enumerate(headers_bva, start=1):
        c = ws_summary.cell(18, c_idx, h)
        c.font = font_tbl_hdr
        c.fill = fill_hdr
        c.alignment = Alignment(horizontal="center")

    bva_data = [
        (
            "4100",
            "Revenue - Product Sales",
            "BLOCKED",
            0.00,
            "BLOCKED",
            "BLOCKED",
            "0 rows committed",
        ),
        (
            "5100",
            "COGS - Materials",
            "BLOCKED",
            17115976.00,
            "BLOCKED",
            "BLOCKED",
            "0 rows committed",
        ),
        (
            "TOTAL",
            "Net Total Variance",
            "BLOCKED",
            17115976.00,
            "BLOCKED",
            "BLOCKED",
            "0 rows committed",
        ),
    ]
    for r_idx, row in enumerate(bva_data, start=19):
        for c_idx, val in enumerate(row, start=1):
            cell = ws_summary.cell(r_idx, c_idx, val)
            cell.font = (
                font_blocked
                if val == "BLOCKED"
                else (font_bold if row[0] == "TOTAL" else font_data)
            )
            cell.fill = fill_blocked if val == "BLOCKED" else PatternFill(fill_type=None)
            cell.border = thin_border
            if isinstance(val, float):
                cell.number_format = "#,##0.00"

    # Section 3: Exceptions Summary
    ws_summary.cell(24, 1, "3. Exception Rules Evaluation").font = font_sec
    ws_summary.cell(
        25, 1, "EVALUATION BLOCKED: FactActual rows not committed due to d365_gl imbalance."
    ).font = font_blocked

    # ----------------- SHEET 2: Difference Log -----------------
    ws_diff.merge_cells("A1:G1")
    ws_diff["A1"] = "Difference Classification Log (GATE-13 / Doc 28 §4.3 & §4.6)"
    ws_diff["A1"].font = font_title

    headers_diff = [
        "Diff ID",
        "Area / Component",
        "Observed Value",
        "Expected Baseline",
        "Delta",
        "Classification",
        "Resolution / Action",
    ]
    for c_idx, h in enumerate(headers_diff, start=1):
        c = ws_diff.cell(3, c_idx, h)
        c.font = font_tbl_hdr
        c.fill = fill_hdr
        c.alignment = Alignment(horizontal="center")

    diff_data = [
        (
            "DIFF-001",
            "D365 GL Sample File Debit/Credit Net Imbalance",
            "₹17,944,515,579.33 imbalance",
            "Balanced Trial Balance",
            "₹17,944,515,579.33",
            "(a) Spec / data generator bug",
            "Generator writes single-sided rows and excludes suspense balancing. Blocks FactActual insertion (IMP-023 fail).",
        ),
        (
            "DIFF-002",
            "FactActual Commitment",
            "0 rows committed",
            "250,037 rows committed",
            "-250,037 rows",
            "(d) Expected difference (safety gate)",
            "Batches rejected upon IMP-023 fail to prevent corruption.",
        ),
        (
            "DIFF-003",
            "Data Quality Score",
            "84",
            "100.0",
            "-16",
            "(a) Spec bug (fixed via DEF-009)",
            "Fix applied in code; score now correctly reflects IMP-023 failure.",
        ),
    ]
    for r_idx, row in enumerate(diff_data, start=4):
        for c_idx, val in enumerate(row, start=1):
            cell = ws_diff.cell(r_idx, c_idx, val)
            cell.font = font_data
            cell.border = thin_border
            cell.alignment = Alignment(wrap_text=True)

    # ----------------- SHEET 3: Sign-Off -----------------
    ws_signoff.merge_cells("A1:F1")
    ws_signoff["A1"] = "GATE-13 Pilot Gate Sign-Off & Approvals"
    ws_signoff["A1"].font = font_title

    headers_sign = [
        "Role",
        "Name",
        "Designation",
        "Date",
        "Status",
        "Notes / Limitation Confirmation",
    ]
    for c_idx, h in enumerate(headers_sign, start=1):
        c = ws_signoff.cell(3, c_idx, h)
        c.font = font_tbl_hdr
        c.fill = fill_hdr
        c.alignment = Alignment(horizontal="center")

    sign_data = [
        (
            "Lead FP&A Analyst",
            "______",
            "Senior Financial Analyst",
            "______",
            "PENDING",
            "To be signed at pilot acceptance.",
        ),
        (
            "Lead Implementation Consultant",
            "______",
            "Lead Implementation Consultant",
            "______",
            "PENDING",
            "To be signed at pilot acceptance.",
        ),
        (
            "Project Owner",
            "______",
            "Project Owner / Release Manager",
            "______",
            "PENDING",
            "To be signed at pilot acceptance.",
        ),
    ]
    for r_idx, row in enumerate(sign_data, start=4):
        for c_idx, val in enumerate(row, start=1):
            cell = ws_signoff.cell(r_idx, c_idx, val)
            cell.font = font_data
            cell.border = thin_border

    # Set column widths
    for ws in [ws_summary, ws_diff, ws_signoff]:
        for col in ws.columns:
            max_len = 0
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if "\n" in val:
                    val = max(val.split("\n"), key=len)
                max_len = max(max_len, len(val))
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    ws_summary.column_dimensions["A"].width = 24
    ws_summary.column_dimensions["B"].width = 38
    ws_diff.column_dimensions["G"].width = 45

    output_path = "packaging/pilot_tieout_worksheet_completed.xlsx"
    wb.save(output_path)
    print(f"Saved completed tie-out worksheet to {output_path}")


if __name__ == "__main__":
    create_completed_tieout()
