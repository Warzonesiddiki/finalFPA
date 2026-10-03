"""
Generate Pilot Tie-Out Worksheet & Difference Classification Log Template (GATE-13 / Doc 28)
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_tieout_template():
    wb = openpyxl.Workbook()
    
    # Setup styles
    font_title = Font(name="Calibri", size=16, bold=True, color="1E293B")
    font_subtitle = Font(name="Calibri", size=11, italic=True, color="64748B")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=11, bold=True, color="0F172A")
    font_normal = Font(name="Calibri", size=11, color="334155")
    
    fill_header = PatternFill(fill_type="solid", start_color="0284C7", end_color="0284C7")
    fill_accent = PatternFill(fill_type="solid", start_color="F1F5F9", end_color="F1F5F9")
    
    border_thin = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    
    # ---------------------------------------------------------
    # Sheet 1: Tie-Out Summary
    # ---------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Tie-Out Summary"
    ws1.views.sheetView[0].showGridLines = True
    
    ws1.append(["FP&A Month-End Pilot Tie-Out Worksheet (GATE-13)"])
    ws1.cell(row=1, column=1).font = font_title
    ws1.append(["Per Doc 28 §4.4 — Comparison of Copilot output against client manual pack"])
    ws1.cell(row=2, column=1).font = font_subtitle
    ws1.append([])
    
    metadata = [
        ("Client / Project Name:", "Northwind Industries — FY26 Real Data Pilot"),
        ("Period Covered:", "2026-P09"),
        ("Build Under Test:", "v0.1.0 — Release Candidate"),
        ("Date of Tie-Out:", "2026-10-02"),
    ]
    for k, v in metadata:
        ws1.append([k, v])
        ws1.cell(row=ws1.max_row, column=1).font = font_bold
        ws1.cell(row=ws1.max_row, column=2).font = font_normal
        
    ws1.append([])
    ws1.append(["Imported Files Summary"])
    ws1.cell(row=ws1.max_row, column=1).font = Font(name="Calibri", size=13, bold=True, color="0F172A")
    ws1.append([])
    
    headers_files = ["File Name", "Source Type", "Batch ID", "Row Count", "Data Quality Score"]
    ws1.append(headers_files)
    for col_idx in range(1, len(headers_files) + 1):
        cell = ws1.cell(row=ws1.max_row, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center")
        
    sample_files = [
        ("d365_gl_actuals.csv", "D365 GL Actuals", "BATCH-001", 12500, "99.2%"),
        ("bank_ledger_actuals.csv", "Bank Ledger", "BATCH-002", 4320, "98.5%"),
        ("payroll_procurement_actuals.csv", "Payroll & Procurement", "BATCH-003", 2850, "100.0%"),
    ]
    for row in sample_files:
        ws1.append(row)
        for col_idx in range(1, len(row) + 1):
            cell = ws1.cell(row=ws1.max_row, column=col_idx)
            cell.font = font_normal
            cell.border = border_thin
            
    ws1.append([])
    ws1.append(["BvA Totals & Key Financial Comparison (MTD / YTD)"])
    ws1.cell(row=ws1.max_row, column=1).font = Font(name="Calibri", size=13, bold=True, color="0F172A")
    ws1.append([])
    
    headers_bva = ["Statement Line / Category", "App Output (MTD)", "Manual Pack (MTD)", "Variance (MTD)", "App Output (YTD)", "Manual Pack (YTD)", "Variance (YTD)"]
    ws1.append(headers_bva)
    for col_idx in range(1, len(headers_bva) + 1):
        cell = ws1.cell(row=ws1.max_row, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center")
        
    bva_rows = [
        ("Revenue", 1250000, 1250000, 0, 11250000, 11250000, 0),
        ("Cost of Goods Sold", 720000, 718500, 1500, 6480000, 6475000, 5000),
        ("Operating Expenses", 310000, 310000, 0, 2790000, 2790000, 0),
        ("Net Operating Income", 220000, 221500, -1500, 1980000, 1985000, -5000),
    ]
    for row in bva_rows:
        ws1.append(row)
        for col_idx in range(1, len(row) + 1):
            cell = ws1.cell(row=ws1.max_row, column=col_idx)
            cell.font = font_normal
            cell.border = border_thin
            if col_idx > 1:
                cell.number_format = '$#,##0'
                cell.alignment = Alignment(horizontal="right")

    # ---------------------------------------------------------
    # Sheet 2: Difference Classification Log
    # ---------------------------------------------------------
    ws2 = wb.create_sheet(title="Difference Classification Log")
    ws2.views.sheetView[0].showGridLines = True
    
    ws2.append(["Difference Classification Log (GATE-13 / Doc 28 §4.3)"])
    ws2.cell(row=1, column=1).font = font_title
    ws2.append(["Taxonomy: (a) Spec bug | (b) Mapping error | (c) Client data / methodology | (d) Expected difference"])
    ws2.cell(row=2, column=1).font = font_subtitle
    ws2.append([])
    
    headers_log = [
        "Diff ID", "Area / Account", "App Amount", "Manual Amount", "Delta", 
        "Classification (a/b/c/d)", "Root Cause Description & Evidence", "Owner of Fix", "Action / Resolution", "Status"
    ]
    ws2.append(headers_log)
    for col_idx in range(1, len(headers_log) + 1):
        cell = ws2.cell(row=ws2.max_row, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center")
        
    sample_diffs = [
        ("DIFF-01", "COGS - Freight & Handling", 720000, 718500, 1500, "(b) Mapping error", "Account 5410 mapped to logistics instead of direct shipping.", "Consultant", "Updated mapping profile mapping rule v1.2", "Resolved"),
        ("DIFF-02", "Depreciation Expense", 45000, 45000, 0, "(d) Expected difference", "Rounding difference of $0.12 normalized by display rules.", "—", "Citing Doc 02 Section 4 rounding spec", "Closed"),
    ]
    for row in sample_diffs:
        ws2.append(row)
        for col_idx in range(1, len(row) + 1):
            cell = ws2.cell(row=ws2.max_row, column=col_idx)
            cell.font = font_normal
            cell.border = border_thin

    # ---------------------------------------------------------
    # Sheet 3: Sign-Off & Approvals
    # ---------------------------------------------------------
    ws3 = wb.create_sheet(title="Sign-Off & Approvals")
    ws3.views.sheetView[0].showGridLines = True
    
    ws3.append(["GATE-13 Real-Data Pilot Sign-Off & Approvals"])
    ws3.cell(row=1, column=1).font = font_title
    ws3.append(["Formal sign-off by FP&A Analyst and Project Owner per Doc 28 §4.5"])
    ws3.cell(row=2, column=1).font = font_subtitle
    ws3.append([])
    
    signoffs = [
        ("Role", "Name", "Signature", "Date", "Approval Status"),
        ("Lead FP&A Analyst", "Priya Sharma", "____________________", "2026-10-02", "APPROVED"),
        ("Client Finance Owner", "Robert Fox", "____________________", "2026-10-02", "APPROVED"),
        ("Implementation Consultant", "Aion Copilot Lead", "____________________", "2026-10-02", "APPROVED"),
    ]
    for row in signoffs:
        ws3.append(row)
        for col_idx in range(1, len(row) + 1):
            cell = ws3.cell(row=ws3.max_row, column=col_idx)
            cell.font = font_header if ws3.max_row == 4 else font_normal
            if ws3.max_row == 4:
                cell.fill = fill_header
            else:
                cell.border = border_thin

    # Auto-fit column widths across sheets
    for sheet in [ws1, ws2, ws3]:
        for col in sheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output_path = "sample-data/templates/pilot_tieout_worksheet_template.xlsx"
    wb.save(output_path)
    print(f"Successfully generated {output_path}")

if __name__ == "__main__":
    create_tieout_template()
