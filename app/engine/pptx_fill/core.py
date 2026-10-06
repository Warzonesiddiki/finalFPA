# Adapted from https://github.com/m3dev/pptx-template @ dc448cb1203443e503a558846dc1ce4c4b7ada3d (Apache-2.0) — ADP-001 —
"""Name-based template resolution and slide assembly.

Copied and edited from ``pptx_template/core.py`` and ``pptx_template/pptx_util.py``
of the pinned upstream commit. The upstream engine binds content to slides with an
expression syntax (``{placeholder}`` tokens) evaluated against a caller-supplied
model; this project binds content by *shape name* instead (docs/12 §3.6), because
the contract names every shape (``PPT-00N_<role>``) and requires a missing name to
abort with ERR-EXP-014 rather than silently skip.

Divergences from upstream (see docs/32 REUSE_AND_PROVENANCE.md, row ``ADP-001``):

* expression-model replacement is dropped; ``resolve_shape`` looks a name up on the
  slide first, then on its layout, and raises :class:`TemplateShapeError` when absent;
* ``edit_slide`` gains :func:`promote_layout_shapes`, which copies the layout's shapes
  onto the slide so each deck slide owns editable shapes instead of rendering inherited
  layout content that cannot be clicked in PowerPoint;
* the upstream expression evaluator (``pyel``), CSV/TSV loading, the Excel model and the
  CLI surface are not carried over — this package fills charts from plain sequences and
  tables from plain matrices (docs/09 ADR-001 runtime deps only);
* upstream log messages and comments were rewritten in this project's own words.

``remove_slide`` and ``set_value_axis`` keep the upstream implementations, which reach
into python-pptx internals that the public API does not expose.
"""

from __future__ import annotations

import copy
from collections.abc import Iterable
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.chart.axis import ValueAxis
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.presentation import Presentation as PresentationObject

#: Template file preferred by the resolver (docs/12 §6 — client base deck).
PRIMARY_TEMPLATE_NAME = "board_pack_template.pptx"
#: Built-in template shipped with the app; the fallback when no base deck is configured.
FALLBACK_TEMPLATE_NAME = "FPAMonthEndCopilot_v1.pptx"

_TEMPLATE_DIR_PARTS = ("app", "engine", "pptx_fill")

_RID_ATTRS = ("r:embed", "r:link", "r:id")
_CHART_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.drawingml.chart+xml"


class TemplateShapeError(Exception):
    """A template layout or one of its named shapes is missing (ERR-EXP-014)."""

    code = "ERR-EXP-014"

    def __init__(self, message: str):
        super().__init__(message)
        self.error_code = self.code


def default_template_dir() -> Path:
    """Return the ``packaging/templates`` directory shipped alongside the ``app`` package."""
    here = Path(__file__).resolve()
    for parent in here.parents:
        candidate = parent / "packaging" / "templates"
        if candidate.is_dir():
            return candidate
    return here.parents[len(_TEMPLATE_DIR_PARTS)] / "packaging" / "templates"


def resolve_template_path(template_dir: Path | None = None) -> Path:
    """Return the template to fill: the base deck when present, else the built-in one.

    Raises:
        TemplateShapeError: when neither file exists (ERR-EXP-014).
    """
    directory = Path(template_dir) if template_dir is not None else default_template_dir()
    primary = directory / PRIMARY_TEMPLATE_NAME
    if primary.exists():
        return primary
    fallback = directory / FALLBACK_TEMPLATE_NAME
    if fallback.exists():
        return fallback
    raise TemplateShapeError(
        f"ERR-EXP-014: the deck template is missing or damaged — neither "
        f"{primary} nor {fallback} exists"
    )


def open_template(template_dir: Path | None = None) -> PresentationObject:
    """Open the resolved template for filling."""
    return Presentation(str(resolve_template_path(template_dir)))


def resolve_layout(prs: PresentationObject, layout_name: str) -> Any:
    """Return the layout named ``layout_name``, or raise ERR-EXP-014."""
    for layout in prs.slide_layouts:
        if layout.name == layout_name:
            return layout
    raise TemplateShapeError(f"ERR-EXP-014: the template is missing layout '{layout_name}'")


def _relocate_element_rels(slide: Any, element: Any, source_part: Any) -> None:
    """Re-point relationship ids in a copied element at the *slide's* part.

    A graphic frame promoted from a layout still references the chart part by the
    relationship id held by the **layout**, which the slide does not have. Leaving the
    id untouched yields a slide that points at a relationship it does not own, so every
    embedded part is re-related to the slide and its id rewritten in place.
    """
    for node in element.iter():
        for attr_name in _RID_ATTRS:
            attr = qn(attr_name)
            old_rid = node.get(attr)
            if not old_rid:
                continue
            rel = source_part.rels.get(old_rid)
            if rel is None:
                continue
            if rel.is_external:
                new_rid = slide.part.relate_to(rel.target_ref, RT.HYPERLINK, is_external=True)
            elif rel.reltype == RT.CHART or rel.target_part.content_type == _CHART_CONTENT_TYPE:
                new_rid = slide.part.relate_to(rel.target_part, RT.CHART)
            else:
                new_rid = slide.part.relate_to(rel.target_part, rel.reltype)
            node.set(attr, new_rid)


def promote_layout_shapes(slide: Any) -> Any:
    """Copy the layout's shapes onto ``slide`` and return it.

    The template ships layouts only — no slides and no placeholders — so a slide added
    for a layout would otherwise show the layout's shapes as inherited content, which
    PowerPoint does not let an editor click. Copying each shape's element onto the slide
    turns them into ordinary per-slide shapes that keep the template's geometry, styling
    and z-order, so the engine fills the template instead of building slides from scratch
    (docs/12 §3.6).
    """
    sp_tree = slide.shapes._spTree
    source_part = slide.slide_layout.part
    for shape in list(slide.slide_layout.shapes):
        element = copy.deepcopy(shape._element)
        _relocate_element_rels(slide, element, source_part)
        sp_tree.append(element)
    return slide


def add_slide_from_layout(prs: PresentationObject, layout_name: str) -> Any:
    """Add a slide for layout ``layout_name`` and promote its shapes onto it."""
    slide = prs.slides.add_slide(resolve_layout(prs, layout_name))
    return promote_layout_shapes(slide)


def resolve_shape(slide: Any, shape_name: str) -> Any:
    """Return the named shape from the slide, falling back to its layout.

    Raises:
        TemplateShapeError: when the name is on neither (ERR-EXP-014).
    """
    for shape in slide.shapes:
        if shape.name == shape_name:
            return shape
    for shape in slide.slide_layout.shapes:
        if shape.name == shape_name:
            return shape
    raise TemplateShapeError(
        f"ERR-EXP-014: missing shape '{shape_name}' in template layout '{slide.slide_layout.name}'"
    )


def has_shape(slide: Any, shape_name: str) -> bool:
    """True when ``shape_name`` resolves on the slide or its layout."""
    try:
        resolve_shape(slide, shape_name)
    except TemplateShapeError:
        return False
    return True


def shape_names(slide: Any) -> list[str]:
    """Every shape name available on the slide plus its layout, in order, de-duplicated."""
    names: list[str] = []
    for source in (slide.shapes, slide.slide_layout.shapes):
        for shape in source:
            if shape.name not in names:
                names.append(shape.name)
    return names


def remove_shape(slide: Any, shape_name: str) -> bool:
    """Delete a named shape from the slide. Returns True when one was removed."""
    for shape in list(slide.shapes):
        if shape.name == shape_name:
            shape._element.getparent().remove(shape._element)
            return True
    return False


def remove_slide(prs: PresentationObject, slide: Any) -> None:
    """Delete ``slide`` from ``prs``, dropping its relationship first."""
    index = next(i for i, sld in enumerate(prs.slides._sldIdLst) if sld.id == slide.slide_id)
    prs.part.drop_rel(prs.slides._sldIdLst[index].rId)
    del prs.slides._sldIdLst[index]


def set_value_axis(
    chart: Any,
    maximum: float | None = None,
    minimum: float | None = None,
    is_second_axis: bool = False,
) -> None:
    """Pin the value axis scale of ``chart`` where the caller has computed bounds."""
    axis = ValueAxis(chart._chartSpace.valAx_lst[1 if is_second_axis else 0])  # type: ignore[no-untyped-call]  # reason: pptx ships no stubs; boundary call
    if maximum is not None:
        axis.maximum_scale = float(maximum)
    if minimum is not None:
        axis.minimum_scale = float(minimum)


def layout_shape_names(prs: PresentationObject, layout_name: str) -> Iterable[str]:
    """Names of every shape on the named layout — the contract a layout must satisfy."""
    return (shape.name for shape in resolve_layout(prs, layout_name).shapes)
