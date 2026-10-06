# Adapted from https://github.com/keithmcnulty/ppt-generation @ 062c4920da96c98574bad0c6cdfd4d10eaa21e02 (CC0-1.0) — ADP-002 —
"""Fill verbs for writing data into a shape that already exists in the template.

Copied and edited from ``edit_pres.py`` of the pinned upstream commit, which demonstrates
the three verbs a deck filler needs: put text in a shape, push series into a native chart
with ``chart.replace_data``, and write cells into a native table. Upstream reaches those
verbs through positional indexing and a pandas dataframe; here every verb takes a named
shape (docs/12 §3.6) and plain Python sequences, so the module has no dependency beyond
``python-pptx`` (docs/09 ADR-001).

Two behaviours differ from the obvious python-pptx calls, and both are load-bearing:

* writing text assigns to an **existing run** rather than to ``text_frame.text`` — the
  latter rebuilds the run and drops its ``rPr``, discarding the font size, weight, colour
  and typeface the template author set;
* filling a table writes into the first run of each cell and clears the rest, so the
  cell keeps the column alignment and font the template carries.

See docs/32 row ``ADP-002`` for the full divergence list.
"""

from __future__ import annotations

import copy
from collections.abc import Sequence
from typing import Any

from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Pt

_TEXT_SHAPE_TYPES = (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.TEXT_BOX)


class TextStyle:
    """Optional run overrides applied on top of the template's own formatting.

    A caller only needs this where the same shape takes different emphasis by data —
    an unfavourable KPI value in red, a favourable one in green — or where a paragraph
    inside a block needs different alignment from the block.
    """

    __slots__ = ("size_pt", "bold", "color", "align")

    def __init__(
        self,
        size_pt: float | None = None,
        bold: bool | None = None,
        color: RGBColor | None = None,
        align: PP_ALIGN | None = None,
    ) -> None:
        self.size_pt = size_pt
        self.bold = bold
        self.color = color
        self.align = align


def select_text_shapes(slide: Any) -> list[Any]:
    """Shapes on ``slide`` that can carry text (auto shapes and text boxes)."""
    return [s for s in slide.shapes if s.shape_type in _TEXT_SHAPE_TYPES]


def select_tables(slide: Any) -> list[Any]:
    """Graphic frames on ``slide`` that carry a table."""
    return [s for s in slide.shapes if s.has_table]


def select_charts(slide: Any) -> list[Any]:
    """Graphic frames on ``slide`` that carry a chart."""
    return [s for s in slide.shapes if s.has_chart]


def _apply_style(run: Any, paragraph: Any, style: TextStyle | None) -> None:
    if style is None:
        return
    if style.size_pt is not None:
        run.font.size = Pt(style.size_pt)
    if style.bold is not None:
        run.font.bold = style.bold
    if style.color is not None:
        run.font.color.rgb = style.color
    if style.align is not None:
        paragraph.alignment = style.align


def set_text(
    shape: Any,
    lines: Any,
    style: TextStyle | None = None,
    styles: Sequence[TextStyle | None] | None = None,
) -> Any:
    """Fill ``shape`` with one or more lines of text, keeping the template's formatting.

    Args:
        shape: A shape carrying a text frame, a table cell, or the text frame itself
            (a chart title is a text frame).
        lines: A single string, or a sequence of strings — one per paragraph.
        style: Overrides applied to every line written.
        styles: Per-line overrides, index-aligned with ``lines``; ``None`` entries fall
            back to ``style``.

    Returns:
        The object passed in, so calls can be chained.
    """
    if isinstance(lines, str):
        lines = [lines]
    text_frame = getattr(shape, "text_frame", shape)
    text_frame.word_wrap = True

    paragraphs = text_frame.paragraphs
    # Reuse the paragraphs the template already carries, cloning the first paragraph's
    # element when more are needed so the extra ones inherit its run properties.
    while len(paragraphs) < len(lines):
        paragraphs[-1]._p.addnext(copy.deepcopy(paragraphs[0]._p))
        paragraphs = text_frame.paragraphs
    for extra in paragraphs[len(lines) :]:
        extra._p.getparent().remove(extra._p)

    for index, line in enumerate(lines):
        paragraph = text_frame.paragraphs[index]
        runs = paragraph.runs
        if not runs:
            paragraph.add_run()
            runs = paragraph.runs
        line_style = styles[index] if styles and index < len(styles) else style
        runs[0].text = line
        _apply_style(runs[0], paragraph, line_style)
        for surplus in runs[1:]:
            surplus.text = ""
    return shape


def set_cell_text(
    cell: Any,
    text: str,
    style: TextStyle | None = None,
    fill_color: RGBColor | None = None,
) -> None:
    """Fill one table cell, preserving its template font, and optionally set its fill."""
    if fill_color is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill_color
    set_text(cell, text, style=style)


def fill_table(
    shape: Any,
    rows: Sequence[Sequence[Any]],
    header: bool = True,
    header_style: TextStyle | None = None,
    cell_style: TextStyle | None = None,
    cell_styles: dict[tuple[int, int], TextStyle | None] | None = None,
    row_styles: dict[int, TextStyle | None] | None = None,
    row_fills: dict[int, RGBColor | None] | None = None,
) -> Any:
    """Write a matrix of values into a table shape that already exists in the template.

    The table keeps the row count, column count, column widths and per-cell formatting it
    was authored with; only text and the requested fills change. Cells outside the matrix
    are left as the template left them.

    Args:
        shape: A shape carrying a table.
        rows: Row-major values, at most ``len(table.rows)`` rows of
            ``len(table.columns)`` cells.
        header: Style row 0 with ``header_style`` instead of ``cell_style``.
        cell_styles: Per-``(row, column)`` overrides.
        row_styles: Per-row overrides.
        row_fills: Per-row background fills.

    Returns:
        The table object.
    """
    table = shape.table
    cell_styles = cell_styles or {}
    row_styles = row_styles or {}
    row_fills = row_fills or {}
    for row_index, values in enumerate(rows):
        if row_index >= len(table.rows):
            break
        for column_index, value in enumerate(values):
            if column_index >= len(table.columns):
                break
            style = cell_styles.get(
                (row_index, column_index),
                row_styles.get(
                    row_index, header_style if (header and row_index == 0) else cell_style
                ),
            )
            set_cell_text(
                table.cell(row_index, column_index),
                "" if value is None else str(value),
                style=style,
                fill_color=row_fills.get(row_index),
            )
    return table


def replace_chart_data(
    shape: Any,
    categories: Sequence[Any],
    series: Sequence[tuple[str, Sequence[Any]]],
    title: str | None = None,
    legend: bool | None = None,
) -> Any:
    """Push new categories and series into a chart the template already contains.

    The chart keeps its type and formatting — series colours, axis settings and legend
    position are authored in the template — and only the data changes. Series must be
    listed in template order; ``None`` leaves a point empty, which is how the gap
    between closed and forecast periods stays visible (docs/12 §5.3).

    Returns:
        The chart object.
    """
    chart = shape.chart
    chart_data = CategoryChartData()  # type: ignore[no-untyped-call]  # reason: pptx ships no stubs; boundary call
    chart_data.categories = ["" if c is None else str(c) for c in categories]
    for name, values in series:
        chart_data.add_series(name, tuple(values))  # type: ignore[no-untyped-call]  # reason: pptx ships no stubs; boundary call
    chart.replace_data(chart_data)

    if title is not None:
        chart.has_title = True
        set_text(chart.chart_title.text_frame, title)
    if legend is not None:
        chart.has_legend = legend
    return chart


def set_notes(slide: Any, text: str) -> Any:
    """Write the speaker notes for a slide, creating the notes slide if needed."""
    slide.notes_slide.notes_text_frame.text = text
    return slide
