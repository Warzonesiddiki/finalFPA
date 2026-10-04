"""File-level ControlTotals worksheet reconciliation (IMP-025)."""

from __future__ import annotations

from datetime import UTC, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

try:
    import openpyxl
except ImportError:  # pragma: no cover - openpyxl is a declared runtime dependency
    openpyxl = None

from app.engine.calc import ZERO, quantize_money
from app.engine.imports.models import ParsedTransaction, ValidationCheckReport
from app.engine.imports.parser import parse_money_value
from app.engine.imports.profiles import normalize_header

HEADER_ALIASES = {
    "scope": {"scope", "controltotalscope", "control_total_scope"},
    "measure": {"measure", "metric", "amounttype", "totaltype"},
    "supplied_total": {
        "suppliedtotal",
        "supplied_total",
        "clienttotal",
        "controltotal",
        "expectedtotal",
        "sourcecontroltotal",
    },
    "tolerance": {"tolerance", "controltotaltolerance", "control_total_tolerance"},
}

MEASURE_ALIASES = {
    "debit": "debit",
    "debits": "debit",
    "totaldebit": "debit",
    "debit_total": "debit",
    "credit": "credit",
    "credits": "credit",
    "totalcredit": "credit",
    "credit_total": "credit",
    "net": "net",
    "netamount": "net",
    "net_amount": "net",
    "nettotal": "net",
    "net_total": "net",
}


def control_totals_not_supplied_report() -> ValidationCheckReport:
    """Return the mandatory, visible IMP-025 skipped result for CSV/no-sheet imports."""
    return ValidationCheckReport(
        check_code="IMP-025",
        check_name="Control-total variance within tolerance",
        status="skipped",
        severity="high",
        offending_count=0,
        skip_reason="no control-totals block supplied",
        detail="No ControlTotals worksheet was supplied.",
    )


def _header_indices(headers: list[Any]) -> dict[str, int]:
    normalized = [normalize_header(str(value or "")).replace(" ", "") for value in headers]
    indices: dict[str, int] = {}
    for canonical, aliases in HEADER_ALIASES.items():
        for index, header in enumerate(normalized):
            if header in aliases:
                indices[canonical] = index
                break
    return indices


def _cell(row: tuple[Any, ...], index: int | None) -> Any:
    return row[index] if index is not None and index < len(row) else None


def _canonical_measure(value: Any, scope: str) -> str:
    raw = normalize_header(str(value or "")).replace(" ", "")
    if not raw:
        raw = normalize_header(scope).replace(" ", "")
    try:
        return MEASURE_ALIASES[raw]
    except KeyError as exc:
        raise ValueError(
            "Measure must identify one file-level total: debit, credit, or net"
        ) from exc


def _loaded_total(measure: str, transactions: list[ParsedTransaction]) -> Decimal:
    field = {
        "debit": "debit",
        "credit": "credit",
        "net": "net_amount",
    }[measure]
    return quantize_money(
        sum((getattr(transaction, field) for transaction in transactions), ZERO)
    )


def read_control_totals_report(
    filepath: str | Path,
    transactions: list[ParsedTransaction],
    *,
    acceptance: dict[str, str] | None = None,
) -> ValidationCheckReport:
    """Read the optional `ControlTotals` worksheet and reconcile supplied totals.

    Workbook contract: row 1 has `Scope`, `Measure`, `SuppliedTotal`, and optional
    `Tolerance`; each subsequent non-empty row is one file-level total. `Measure`
    is debit, credit, or net. A supplied acceptance applies to any variance beyond
    tolerance and is copied into the IMP-025 evidence rows before the batch commits.
    """
    if openpyxl is None:  # pragma: no cover
        raise ImportError("openpyxl is required to read the ControlTotals worksheet")

    workbook = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    try:
        if "ControlTotals" not in workbook.sheetnames:
            return control_totals_not_supplied_report()

        worksheet = workbook["ControlTotals"]
        if worksheet.sheet_state != "visible":
            return ValidationCheckReport(
                check_code="IMP-025",
                check_name="Control-total variance within tolerance",
                status="fail",
                severity="high",
                offending_count=1,
                detail="ControlTotals worksheet is hidden and cannot be used as import evidence.",
                message_slug="import.controlTotalVariance",
            )
        row_iter = worksheet.iter_rows(values_only=True)
        header_row = next(row_iter, None)
        if not header_row:
            return ValidationCheckReport(
                check_code="IMP-025",
                check_name="Control-total variance within tolerance",
                status="fail",
                severity="high",
                offending_count=1,
                detail="ControlTotals worksheet is empty; expected a header row.",
                message_slug="import.controlTotalVariance",
            )

        indices = _header_indices(list(header_row))
        required = {"scope", "supplied_total"}
        missing = sorted(required - indices.keys())
        if missing:
            return ValidationCheckReport(
                check_code="IMP-025",
                check_name="Control-total variance within tolerance",
                status="fail",
                severity="high",
                offending_count=1,
                detail=(
                    "ControlTotals worksheet is missing required column(s): "
                    + ", ".join(missing)
                    + ". Required columns: Scope, SuppliedTotal; Measure is required "
                    "unless Scope itself is debit, credit, or net."
                ),
                message_slug="import.controlTotalVariance",
            )

        totals: list[dict[str, Any]] = []
        errors: list[str] = []
        seen_scopes: set[str] = set()
        variance_count = 0
        accepted_variances = 0
        accepted_by = str(acceptance.get("accepted_by", "")).strip() if acceptance else ""
        acceptance_reason = str(acceptance.get("reason", "")).strip() if acceptance else ""
        if acceptance is not None and (
            not accepted_by or len(acceptance_reason) < 10
        ):
            raise ValueError(
                "Control-total acceptance requires accepted_by and a reason of at least 10 characters"
            )
        accepted_at = datetime.now(UTC).isoformat() if acceptance else None

        for worksheet_row, row in enumerate(row_iter, start=2):
            if not row or all(value is None or str(value).strip() == "" for value in row):
                continue

            scope = str(_cell(row, indices.get("scope")) or "").strip()
            source_row_ref = f"ControlTotals!{worksheet_row}"
            evidence: dict[str, Any] = {
                "scope": scope,
                "source_row_ref": source_row_ref,
                "accepted": False,
            }
            try:
                if not scope:
                    raise ValueError("Scope is required")
                if scope.casefold() in seen_scopes:
                    raise ValueError(f"Duplicate scope '{scope}'")
                seen_scopes.add(scope.casefold())

                measure = _canonical_measure(
                    _cell(row, indices.get("measure")), scope
                )
                supplied_value = _cell(row, indices.get("supplied_total"))
                if supplied_value in (None, ""):
                    raise ValueError("SuppliedTotal is required")
                supplied = parse_money_value(supplied_value)
                tolerance_value = _cell(row, indices.get("tolerance"))
                tolerance = (
                    parse_money_value(tolerance_value)
                    if tolerance_value not in (None, "")
                    else ZERO
                )
                if tolerance < ZERO:
                    raise ValueError("Tolerance cannot be negative")

                loaded = _loaded_total(measure, transactions)
                variance = quantize_money(loaded - supplied)
                beyond_tolerance = abs(variance) > tolerance
                accepted = beyond_tolerance and bool(acceptance)
                evidence.update(
                    {
                        "measure": measure,
                        "supplied_total": str(supplied),
                        "loaded_total": str(loaded),
                        "variance": str(variance),
                        "tolerance": str(tolerance),
                        "accepted": accepted,
                    }
                )
                if beyond_tolerance:
                    variance_count += 1
                    if accepted:
                        accepted_variances += 1
                        evidence.update(
                            {
                                "accepted_by": accepted_by,
                                "acceptance_reason": acceptance_reason,
                                "accepted_at": accepted_at,
                            }
                        )
            except (InvalidOperation, ValueError, TypeError) as exc:
                evidence["error"] = str(exc)
                errors.append(f"{source_row_ref}: {exc}")
            totals.append(evidence)

        if not totals:
            return ValidationCheckReport(
                check_code="IMP-025",
                check_name="Control-total variance within tolerance",
                status="fail",
                severity="high",
                offending_count=1,
                detail="ControlTotals worksheet has no control-total rows.",
                message_slug="import.controlTotalVariance",
            )

        if errors or variance_count > accepted_variances:
            status = "fail"
        elif accepted_variances:
            status = "warn"
        else:
            status = "pass"

        detail = (
            f"Compared {len(totals)} file-level control total(s); "
            f"{variance_count} exceeded tolerance and {accepted_variances} "
            "were explicitly accepted."
        )
        if errors:
            detail += " Errors: " + "; ".join(errors)
        elif variance_count > accepted_variances:
            detail += " Import is blocked until each variance is within tolerance or accepted."
        return ValidationCheckReport(
            check_code="IMP-025",
            check_name="Control-total variance within tolerance",
            status=status,
            severity="high",
            offending_count=variance_count + len(errors),
            sample_rows=totals,
            detail=detail,
            message_slug="import.controlTotalVariance" if status in {"fail", "warn"} else None,
        )
    finally:
        workbook.close()
