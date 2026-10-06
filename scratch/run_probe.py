import copy
import sys

sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation  # noqa: E402

p = Presentation("packaging/templates/FPAMonthEndCopilot_v1.pptx")
lay = next(l for l in p.slide_layouts if l.name == "FPA-PPT-002")
s = p.slides.add_slide(lay)
for sh in list(s.slide_layout.shapes):
    s.shapes._spTree.append(copy.deepcopy(sh._element))


def by_name(slide, n):
    return next(x for x in slide.shapes if x.name == n)


def report(nm, sh):
    tf = sh.text_frame
    for para in tf.paragraphs:
        for r in para.runs:
            print(f"  {nm}: run.text={r.text!r} sz={r.font.size} bold={r.font.bold}")
            try:
                print(f"  {nm}: rgb={r.font.color.rgb}")
            except Exception as e:
                print(f"  {nm}: rgb err {e}")


print("--- A: run.text assignment (preserve) ---")
t = by_name(s, "PPT-002_title")
t.text_frame.paragraphs[0].runs[0].text = "Via run.text"
report("title", t)
print(t._element.xml[:600])

print("--- B: text_frame.text assignment (destroys) ---")
k = by_name(s, "PPT-002_kicker")
k.text_frame.text = "Via tf.text"
report("kicker", k)
print(k._element.xml[:600])

print("--- C: multi-line via add_paragraph then clone rPr ---")
n = by_name(s, "PPT-002_narrative")
tf = n.text_frame
base = tf.paragraphs[0].runs[0]._r
print("narrative base rPr present:", base.find("{http://schemas.openxmlformats.org/drawingml/2006/main}rPr") is not None)