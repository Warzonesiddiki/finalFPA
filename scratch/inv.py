import sys

sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation  # noqa: E402
from pptx.util import Emu  # noqa: E402

p = Presentation("packaging/templates/FPAMonthEndCopilot_v1.pptx")
print(
    "slide size EMU:", p.slide_width, p.slide_height,
    "inches:", round(Emu(p.slide_width).inches, 4), round(Emu(p.slide_height).inches, 4),
)
for i, lay in enumerate(p.slide_layouts):
    print(f"\n=== LAYOUT[{i}] name={lay.name!r} shapes={len(lay.shapes)}")
    for sh in lay.shapes:
        ph = ""
        if sh.is_placeholder:
            ph = f" PH[idx={sh.placeholder_format.idx} type={sh.placeholder_format.type}]"
        t = ""
        if sh.has_text_frame:
            t = " text=" + repr(sh.text_frame.text[:44])
        extra = ""
        if sh.has_chart:
            ch = sh.chart
            extra = f" chart[type={ch.chart_type} series={len(ch.plots[0].series)} cats={len(list(ch.plots[0].categories))}]"
        if sh.has_table:
            tb = sh.table
            extra = f" table[{len(tb.rows)}x{len(tb.columns)}] hdr={tb.cell(0,0).text!r}"
        print(f"  id={sh.shape_id:>4} {sh.name!r} type={sh.shape_type}{ph}{extra}{t}")
print("\nslides in template:", len(p.slides._sldIdLst))
