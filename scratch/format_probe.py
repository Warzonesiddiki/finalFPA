import sys

sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation  # noqa: E402

p = Presentation("packaging/templates/FPAMonthEndCopilot_v1.pptx")
lay = next(l for l in p.slide_layouts if l.name == "FPA-PPT-002")

for nm in ("PPT-002_title", "PPT-002_kicker", "PPT-002_narrative", "PPT-002_kpi1_value", "PPT-002_narrative_label"):
    sh = next(s for s in lay.shapes if s.name == nm)
    xml = sh._element.xml
    print("=" * 70)
    print(nm)
    print(xml[:1400])
print("=" * 70)
print("### now: write text into promoted copies and inspect inherited formatting")

import copy  # noqa: E402

s = p.slides.add_slide(lay)
for sh in list(s.slide_layout.shapes):
    s.shapes._spTree.append(copy.deepcopy(sh._element))


def by_name(slide, n):
    return next(x for x in slide.shapes if x.name == n)


t = by_name(s, "PPT-002_title")
t.text_frame.text = "Inherited?"
for para in t.text_frame.paragraphs:
    for r in para.runs:
        print("title run size:", r.font.size, "bold:", r.font.bold, "color type:", r.font.color.type)
        try:
            print("title run rgb:", r.font.color.rgb)
        except Exception as e:
            print("title rgb err:", e)
print("title xml after write:", t._element.xml[:900])