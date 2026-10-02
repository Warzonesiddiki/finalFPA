#!/usr/bin/env python3
"""Evidence-backed Phase 0 documentation re-audit.

This is a governance checker, not product code.  It validates documentary evidence for
Kickoff + Addons 1–5 and deliberately does not claim that unimplemented product
features have been executed.  It emits one result for each of the 70 Phase-0 gate
checks defined in docs/14_TESTING_QA_PLAN.md §15.

Run from the repository root:
    python audit/phase0_reaudit.py --report evidence/<date>-phase0-re-audit/phase0_reaudit_report.md
"""
from __future__ import annotations

import argparse
import hashlib
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OFFICIAL_ADDON_5 = ROOT / "project prompt" / "ADDON_5_OWNER_CONTROL_EVIDENCE_DATA_EGRESS_EXTENSION.md"
ORACLE = ROOT / "tests" / "oracle" / "golden_month_hand_check_template.xlsx"
EXPECTED_ADDON_5_SHA256 = "cfbb69411d586194d2ad8ef6034a74e19bcb0ae976b466485adda75a97372b11"


@dataclass(frozen=True)
class Result:
    check_id: str
    source: str
    assertion: str
    evidence: str
    passed: bool
    detail: str


class Auditor:
    def __init__(self) -> None:
        self.results: list[Result] = []

    def add(self, check_id: str, source: str, assertion: str, evidence: str, condition: bool, detail: str) -> None:
        self.results.append(Result(check_id, source, assertion, evidence, bool(condition), detail))


def text(relative: str) -> str:
    path = ROOT / relative
    if not path.is_file():
        return ""
    return path.read_text(encoding="utf-8")


def has(relative: str, *needles: str) -> bool:
    body = text(relative).casefold()
    return bool(body) and all(needle.casefold() in body for needle in needles)


def exists(relative: str) -> bool:
    return (ROOT / relative).exists()


def section_between(body: str, start: str, end: str) -> str:
    try:
        tail = body.split(start, 1)[1]
    except IndexError:
        return ""
    return tail.split(end, 1)[0] if end in tail else tail


def coverage_rows(prefix: str, start: str, end: str) -> list[str]:
    body = section_between(text("docs/00_INDEX.md"), start, end)
    return [line for line in body.splitlines() if line.startswith(f"| {prefix}")]


def local_markdown_links() -> tuple[bool, str]:
    """Check inline repository-local Markdown links, excluding URLs and anchors."""
    pattern = re.compile(r"(?<!!)(?:\[[^\]]*\])\(([^)\s]+)(?:\s+[^)]*)?\)")
    failures: list[str] = []
    files = sorted(ROOT.rglob("*.md"))
    for file in files:
        if any(part in {".git", "node_modules"} for part in file.parts):
            continue
        for line_no, line in enumerate(file.read_text(encoding="utf-8").splitlines(), 1):
            for raw_target in pattern.findall(line):
                target = raw_target.strip("<>")
                if target.startswith(("#", "http://", "https://", "mailto:", "tel:")):
                    continue
                target_path = target.split("#", 1)[0]
                if not target_path:
                    continue
                candidate = (file.parent / target_path).resolve()
                try:
                    candidate.relative_to(ROOT)
                except ValueError:
                    failures.append(f"{file.relative_to(ROOT)}:{line_no}: escapes repository: {target}")
                    continue
                if not candidate.exists():
                    failures.append(f"{file.relative_to(ROOT)}:{line_no}: missing: {target}")
    if failures:
        return False, "; ".join(failures[:8])
    return True, f"checked {len(files)} Markdown files; all repository-local inline links resolve"


def oracle_workbook_check() -> tuple[bool, str]:
    if not ORACLE.is_file():
        return False, "oracle workbook is missing"
    required_sheets = {"Hand Check", "Raw actuals", "Budget rows", "Forecast rows"}
    try:
        import openpyxl  # type: ignore[import-not-found]

        workbook = openpyxl.load_workbook(ORACLE, data_only=False, read_only=False)
        sheets = set(workbook.sheetnames)
        formulas = [
            cell.value
            for worksheet in workbook.worksheets
            for row in worksheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        ]
        workbook.close()
    except Exception as exc:  # pragma: no cover - records exact validation error
        return False, f"openpyxl could not load workbook: {exc!r}"
    with ZipFile(ORACLE) as archive:
        bad = archive.testzip()
    required_formulas = ("SUMIFS", "IF(", "B9-B10")
    passed = not bad and required_sheets.issubset(sheets) and all(
        any(token in formula.upper() for formula in formulas) for token in required_formulas
    )
    detail = (
        f"sheets={sorted(sheets)}; formulas={len(formulas)}; "
        f"zip_test={'OK' if not bad else bad}; openpyxl={openpyxl.__version__}"
    )
    return passed, detail


def header_check() -> tuple[bool, str]:
    failures: list[str] = []
    for number in range(31):
        candidates = list(DOCS.glob(f"{number:02d}_*.md"))
        if len(candidates) != 1:
            failures.append(f"{number:02d}: expected one numbered document, found {len(candidates)}")
            continue
        lines = candidates[0].read_text(encoding="utf-8").splitlines()
        preamble = "\n".join(lines[:20])
        if "> **Status:" not in preamble or "> **TL;DR" not in preamble:
            failures.append(candidates[0].name)
    return not failures, "all 31 numbered docs have Status + TL;DR headers" if not failures else ", ".join(failures)


def no_product_code_check() -> tuple[bool, str]:
    code_suffixes = {".py", ".ts", ".tsx", ".js", ".jsx", ".css", ".html", ".rs", ".go"}
    offenders: list[str] = []
    for directory in (ROOT / "app", ROOT / "ui", ROOT / "packaging", ROOT / "scripts"):
        if not directory.exists():
            continue
        for path in directory.rglob("*"):
            if path.is_file() and path.suffix.casefold() in code_suffixes:
                offenders.append(str(path.relative_to(ROOT)))
    return not offenders, "no product source files found under app/, ui/, packaging/ or scripts/" if not offenders else ", ".join(offenders)


def audit() -> tuple[list[Result], dict[str, tuple[bool, str]]]:
    a = Auditor()
    doc00 = text("docs/00_INDEX.md")

    # GATE-01 — Kickoff §5 (9)
    a.add("GATE-01-01", "Kickoff §5", "FRs are numbered, testable and traced with no blocking TBD.", "02 + 20", has("docs/02_FUNCTIONAL_SPEC.md", "acceptance criteria", "FR-ONB-001") and has("docs/20_REQUIREMENTS_TRACEABILITY.md", "All 156 FRs", "Test ID"), "FR and traceability markers found")
    a.add("GATE-01-02", "Kickoff §5", "Exact calculation examples cover the required finance and every forecast method.", "05 §12 + 07 §11", has("docs/05_CALCULATION_SPEC.md", "F1", "F14", "F14e", "F14f", "F14g") and has("docs/07_FORECAST_METHODS_SPEC.md", "F14e", "F14f", "F14g"), "F1–F14 plus locked/three-month/manual examples found")
    a.add("GATE-01-03", "Kickoff §5", "At least 15 fully specified exception rules and planted tests exist.", "06", has("docs/06_EXCEPTION_RULES_CATALOG.md", "EXC-001", "EXC-024", "planted"), "24-rule and planting markers found")
    a.add("GATE-01-04", "Kickoff §5", "UI spec covers screens and required non-happy states for non-technical users.", "08", has("docs/08_UI_UX_SPEC.md", "SCR-043", "empty", "loading", "first-run", "non-technical"), "screen and state markers found")
    a.add("GATE-01-05", "Kickoff §5", "PowerPoint contract specifies 4–6 editable/native slides and branding.", "12", has("docs/12_POWERPOINT_OUTPUT_SPEC.md", "PPT-001", "PPT-006", "editable", "native", "brand"), "six-slide editable/native contract markers found")
    a.add("GATE-01-06", "Kickoff §5", "A clean-Windows installer script is specified for the post-approval packaging spike.", "15", has("docs/15_PACKAGING_DEPLOYMENT_RUNBOOK.md", "clean Windows 11", "step-by-step", "installer"), "documentation-only readiness check; no installer is claimed built")
    a.add("GATE-01-07", "Kickoff §5", "The required numeric NFRs are stated.", "14 §3", has("docs/14_TESTING_QA_PLAN.md", "NFR-001", "NFR-016", "250k", "500 MB"), "NFR range and core targets found")
    a.add("GATE-01-08", "Kickoff §5", "Unconfirmed items are in the OQ/questionnaire registers.", "18 + 21", has("docs/18_GLOSSARY_ASSUMPTIONS_OPEN_QUESTIONS.md", "OQ-001", "OQ-023", "Register size:") and has("docs/21_CLIENT_ONBOARDING_QUESTIONNAIRE.md", "Q-001", "Q-021"), "OQ and Q registers found")
    a.add("GATE-01-09", "Kickoff §5", "Every FR links to a future test.", "20", has("docs/20_REQUIREMENTS_TRACEABILITY.md", "All 156 FRs", "every FR has"), "traceability total and test column markers found")

    # GATE-02 — Addon 1 §O (12)
    a.add("GATE-02-01", "Addon 1 §O", "Docs 21–25 exist and all questionnaire items carry defaults.", "21–25", all(exists(f"docs/{n:02d}_{name}.md") for n, name in [(21, "CLIENT_ONBOARDING_QUESTIONNAIRE"), (22, "END_USER_GUIDE"), (23, "CONSULTANT_HANDOVER_AND_SUPPORT"), (24, "RELEASE_AND_VERSIONING_RUNBOOK"), (25, "RISK_REGISTER")]) and has("docs/21_CLIENT_ONBOARDING_QUESTIONNAIRE.md", "21 structured questions", "default"), "docs 21–25 and 21-question default markers found")
    a.add("GATE-02-02", "Addon 1 §O", "Addon 1 additions are integrated and recorded in the changelog.", "00 + CHANGELOG", bool(coverage_rows("A1-", "### 4.2", "### 4.3")) and all("INTEGRATED" in row for row in coverage_rows("A1-", "### 4.2", "### 4.3")) and has("docs/CHANGELOG.md", "Addon 1"), "all Addon-1 coverage rows integrated; changelog contains history")
    a.add("GATE-02-03", "Addon 1 §O", "The documented analyst month-end tabletop is recorded.", "evidence table-top record", exists("evidence/2026-10-02-phase0-re-audit/tabletop_walkthroughs.md") and has("evidence/2026-10-02-phase0-re-audit/tabletop_walkthroughs.md", "Analyst month-end", "pack issuance"), "requires current tabletop evidence file")
    a.add("GATE-02-04", "Addon 1 §O", "Excel/CSV hardening maps quirks to handling and error copy.", "04", has("docs/04_SOURCE_MAPPING_AND_IMPORT_SPEC.md", "IMP-032", "handle-or-reject", "ERR-IMP"), "hardening markers found")
    a.add("GATE-02-05", "Addon 1 §O", "OneDrive/storage decision and test are specified.", "09 + 14", has("docs/09_TECHNICAL_ARCHITECTURE.md", "ADR-004", "OneDrive") and has("docs/14_TESTING_QA_PLAN.md", "TST-WIN-06"), "ADR-004 and Windows test found")
    a.add("GATE-02-06", "Addon 1 §O", "SmartScreen/signing decision and non-technical mitigation are specified.", "09", has("docs/09_TECHNICAL_ARCHITECTURE.md", "ADR-003", "SmartScreen", "non-technical"), "ADR-003 mitigation markers found")
    a.add("GATE-02-07", "Addon 1 §O", "Upgrade/migration test names a prior-version fixture.", "24 + 14", has("docs/24_RELEASE_AND_VERSIONING_RUNBOOK.md", "prior-version", "fixture") and has("docs/14_TESTING_QA_PLAN.md", "TST-E2E-05"), "fixture and TST-E2E-05 markers found")
    a.add("GATE-02-08", "Addon 1 §O", "Prompt-injection test case and reproducible synthetic fixture are specified.", "14 + generator", has("docs/14_TESTING_QA_PLAN.md", "TST-SEC-14", "INJ-01") and has("sample-data/generate_sample_data.py", "INJ-01", "EXC-SEC-14"), "test case plus generator fixture found")
    a.add("GATE-02-09", "Addon 1 §O", "Licence allow-list, secret scan, SBOM and licence artefact are specified.", "13 + 15 + 24", has("docs/13_SECURITY_PRIVACY.md", "allow-list", "secret scan") and has("docs/15_PACKAGING_DEPLOYMENT_RUNBOOK.md", "THIRD_PARTY_LICENSES", "SBOM") and has("docs/24_RELEASE_AND_VERSIONING_RUNBOOK.md", "SBOM"), "supply-chain controls found")
    a.add("GATE-02-10", "Addon 1 §O", "Task-structured end-user guide and training/support outline exist.", "22 + 23", has("docs/22_END_USER_GUIDE.md", "T-01", "T-21", "training") and has("docs/23_CONSULTANT_HANDOVER_AND_SUPPORT.md", "incident", "training"), "guide/training markers found")
    a.add("GATE-02-11", "Addon 1 §O", "Parked Addon 1 scope is preserved in the backlog.", "01 + 27", has("docs/27_BACKLOG.md", "BL-016", "BL-036", "promotion") and has("docs/01_PRD.md", "out of scope"), "backlog range and PRD boundary found")
    a.add("GATE-02-12", "Addon 1 §O", "Addon 1 NFR numbers are in the test plan.", "14 §3", has("docs/14_TESTING_QA_PLAN.md", "NFR-001", "NFR-012", "NFR-016"), "NFR markers found")

    # GATE-03 — Addon 2 §I (12)
    a.add("GATE-03-01", "Addon 2 §I", "Coverage Matrix contains integrated Kickoff/Addons 1–2 rows.", "00 §4", len(coverage_rows("K-S", "### 4.1", "### 4.2")) == 15 and len(coverage_rows("A1-", "### 4.2", "### 4.3")) == 17 and len(coverage_rows("A2-", "### 4.3", "### 4.4")) == 17 and all("INTEGRATED" in row for row in coverage_rows("A2-", "### 4.3", "### 4.4")), "expected coverage-row counts and statuses found")
    a.add("GATE-03-02", "Addon 2 §I", "API contract defines OpenAPI/type generation, envelope and pagination.", "26", has("docs/26_API_CONTRACT.md", "OpenAPI", "type-generation", "envelope", "pagination"), "API contract markers found")
    a.add("GATE-03-03", "Addon 2 §I", "Headless engine boundary and CLI exit codes are documented.", "09", has("docs/09_TECHNICAL_ARCHITECTURE.md", "headless", "exit code", "CLI"), "boundary and CLI markers found")
    a.add("GATE-03-04", "Addon 2 §I", "ADR-002 and toolchain traceability are documented.", "09", has("docs/09_TECHNICAL_ARCHITECTURE.md", "ADR-002", "ruff", "pytest", "Playwright"), "ADR/toolchain markers found")
    a.add("GATE-03-05", "Addon 2 §I", "Data-volume/config layering and test cases are documented.", "09 + 14", has("docs/09_TECHNICAL_ARCHITECTURE.md", "configuration layering", "data-volume") and has("docs/14_TESTING_QA_PLAN.md", "TST-API-09", "TST-UI-15"), "architecture and test markers found")
    a.add("GATE-03-06", "Addon 2 §I", "Exception identity/re-run semantics include a scenario.", "06", has("docs/06_EXCEPTION_RULES_CATALOG.md", "re-run", "identity", "worked"), "identity/re-run markers found")
    a.add("GATE-03-07", "Addon 2 §I", "Forecast accuracy and TTM formulas have examples.", "05", has("docs/05_CALCULATION_SPEC.md", "TTM", "forecast accuracy", "F14"), "formula/example markers found")
    a.add("GATE-03-08", "Addon 2 §I", "Screen/API traceability chain exists.", "08 + 20 + 26", has("docs/08_UI_UX_SPEC.md", "SCR-001", "SCR-043") and has("docs/20_REQUIREMENTS_TRACEABILITY.md", "Screen", "API") and has("docs/26_API_CONTRACT.md", "95"), "screen/API chain markers found")
    a.add("GATE-03-09", "Addon 2 §I", "PPT character budgets are per placeholder.", "12", has("docs/12_POWERPOINT_OUTPUT_SPEC.md", "character budget", "placeholder"), "character-budget markers found")
    a.add("GATE-03-10", "Addon 2 §I", "Coverage bars, scripts/check and E2E list are recorded.", "14", has("docs/14_TESTING_QA_PLAN.md", "scripts/check", "coverage", "E2E"), "quality-runner markers found")
    a.add("GATE-03-11", "Addon 2 §I", "PRD decisions and canonical disclaimer are resolved.", "01", has("docs/01_PRD.md", "disclaimer", "DEC-", "scope"), "decision/disclaimer markers found")
    a.add("GATE-03-12", "Addon 2 §I", "AI usage log, provenance and number mismatch stance are specified.", "10", has("docs/10_AI_INTEGRATION_SPEC.md", "usage log", "provenance", "number-mismatch"), "AI-control markers found")

    # GATE-04 — Addon 3 §J (12)
    a.add("GATE-04-01", "Addon 3 §J", "Docs 27/28 exist with backlog and acceptance mechanics.", "27 + 28", exists("docs/27_BACKLOG.md") and exists("docs/28_ACCEPTANCE_UAT_AND_GO_LIVE.md") and has("docs/27_BACKLOG.md", "promotion") and has("docs/28_ACCEPTANCE_UAT_AND_GO_LIVE.md", "sign-off"), "both documents and core markers found")
    a.add("GATE-04-02", "Addon 3 §J", "All Addon 3 Coverage Matrix rows are integrated.", "00 §4.4", len(coverage_rows("A3-", "### 4.4", "### 4.5")) == 17 and all("INTEGRATED" in row for row in coverage_rows("A3-", "### 4.4", "### 4.5")), "17 Addon-3 rows integrated")
    a.add("GATE-04-03", "Addon 3 §J", "Four full prompt texts with worked examples exist.", "10", all(has("docs/10_AI_INTEGRATION_SPEC.md", f"PROMPT-0{number}", "input", "output") for number in range(1, 5)), "PROMPT-01…04 and example markers found")
    a.add("GATE-04-04", "Addon 3 §J", "Mapping review, commentary lock and issuance register are specified.", "02 + 03 + 08", has("docs/02_FUNCTIONAL_SPEC.md", "Mapping Review Queue", "commentary") and has("docs/03_DATA_DICTIONARY.md", "issuance") and has("docs/08_UI_UX_SPEC.md", "lock"), "workflow markers found")
    a.add("GATE-04-05", "Addon 3 §J", "Chart inventory and centralised conditional formatting exist.", "08", has("docs/08_UI_UX_SPEC.md", "Chart inventory", "CF-012"), "chart/CF markers found")
    a.add("GATE-04-06", "Addon 3 §J", "Output conventions and cross-artifact testing are specified.", "11 + 12 + 14", has("docs/11_EXCEL_OUTPUT_SPEC.md", "cross-artifact") and has("docs/12_POWERPOINT_OUTPUT_SPEC.md", "cross-artifact") and has("docs/14_TESTING_QA_PLAN.md", "cross-artifact"), "cross-artifact markers found")
    a.add("GATE-04-07", "Addon 3 §J", "The generator creates the negative-file corpus with 16 cases.", "generator + regeneration evidence", has("sample-data/generate_sample_data.py", "16 malformed", "malformed") and exists("evidence/2026-10-02-phase0-re-audit/sample_data_regeneration.md"), "source plus current regeneration evidence required")
    a.add("GATE-04-08", "Addon 3 §J", "Data-quality score formula has a worked example.", "05", has("docs/05_CALCULATION_SPEC.md", "data-quality score", "worked example"), "DQ formula/example markers found")
    a.add("GATE-04-09", "Addon 3 §J", "PRD settles success metrics, IP/licensing and forced scope choices.", "01", has("docs/01_PRD.md", "success metrics", "IP", "forced"), "PRD decision markers found")
    a.add("GATE-04-10", "Addon 3 §J", "Error-code families and message catalogue rules are specified.", "26 + 08", has("docs/26_API_CONTRACT.md", "error catalogue", "ERR-") and has("docs/08_UI_UX_SPEC.md", "message catalog"), "error/message markers found")
    a.add("GATE-04-11", "Addon 3 §J", "Project DoD, UAT entry/exit, severities, 23-item go-live and sign-off exist.", "28", has("docs/28_ACCEPTANCE_UAT_AND_GO_LIVE.md", "Definition of Done", "entry", "S1", "23-item", "sign-off"), "acceptance markers found")
    a.add("GATE-04-12", "Addon 3 §J", "Decided log and ADR index are active.", "18 + 09", has("docs/18_GLOSSARY_ASSUMPTIONS_OPEN_QUESTIONS.md", "DEC-040") and has("docs/09_TECHNICAL_ARCHITECTURE.md", "ADR-000", "ADR-002"), "decision/ADR markers found")

    # GATE-05 — Addon 4 §K (13)
    a.add("GATE-05-01", "Addon 4 §K", "All Addon 4 Coverage Matrix rows are integrated.", "00 §4.5", len(coverage_rows("A4-", "### 4.5", "### 4.6")) == 12 and all("INTEGRATED" in row for row in coverage_rows("A4-", "### 4.5", "### 4.6")), "12 Addon-4 rows integrated")
    headers_ok, headers_detail = header_check()
    a.add("GATE-05-02", "Addon 4 §K", "Docs 00–30 have standard Status/TL;DR headers.", "docs 00–30", headers_ok, headers_detail)
    a.add("GATE-05-03", "Addon 4 §K", "Source-of-Truth Matrix and conflict rule are documented.", "00 §5–6", has("docs/00_INDEX.md", "Source-of-Truth Matrix", "conflict", "one owner"), "source/conflict markers found")
    a.add("GATE-05-04", "Addon 4 §K", "Quote-before-code and FR citation rules are written.", "19", has("docs/19_VIBE_CODING_PLAYBOOK.md", "quote-before-code", "FR"), "protocol markers found")
    a.add("GATE-05-05", "Addon 4 §K", "Priorities, never-cut list and cut process are documented.", "02 + 16", has("docs/02_FUNCTIONAL_SPEC.md", "P0", "P1", "P2", "never-cut") and has("docs/16_ROADMAP_PHASES.md", "cut"), "priority/cut markers found")
    a.add("GATE-05-06", "Addon 4 §K", "Per-phase estimates are documented.", "16", has("docs/16_ROADMAP_PHASES.md", "ideal developer-days", "estimate"), "estimate markers found")
    a.add("GATE-05-07", "Addon 4 §K", "Client requirements pack is plain-language and has sign-off.", "29", has("docs/29_CLIENT_REQUIREMENTS_PACK.md", "sign-off", "plain"), "client-pack markers found")
    a.add("GATE-05-08", "Addon 4 §K", "Approval recording and post-approval impact rules exist.", "19 + CHANGELOG", has("docs/19_VIBE_CODING_PLAYBOOK.md", "approval", "impact note") and has("docs/CHANGELOG.md", "approval"), "approval/impact markers found")
    a.add("GATE-05-09", "Addon 4 §K", "Real-data pilot/tie-out and classification are defined.", "28", has("docs/28_ACCEPTANCE_UAT_AND_GO_LIVE.md", "tie-out", "classification", "GATE-13"), "pilot markers found")
    a.add("GATE-05-10", "Addon 4 §K", "Tolerance and edge-case matrices with message IDs exist.", "05 + 02 + 14", has("docs/05_CALCULATION_SPEC.md", "tolerance", "minor unit") and has("docs/02_FUNCTIONAL_SPEC.md", "edge-case", "ERR-") and has("docs/14_TESTING_QA_PLAN.md", "edge-case", "ERR-"), "tolerance/matrix markers found")
    a.add("GATE-05-11", "Addon 4 §K", "Sample-data watermark/project type/non-delivery rules are specified.", "generator + 14", has("sample-data/generate_sample_data.py", "WATERMARK", "PROJECT_TYPE") and has("docs/14_TESTING_QA_PLAN.md", "Never delivered", "project_type"), "generator and governance markers found")
    a.add("GATE-05-12", "Addon 4 §K", "Spike/fresh-clone/code-health/storage controls are specified.", "09 + 14 + 17", has("docs/09_TECHNICAL_ARCHITECTURE.md", "spike", "storage") and has("docs/14_TESTING_QA_PLAN.md", "fresh-clone") and has("docs/17_CODING_STANDARDS.md", "code-health"), "engineering-control markers found")
    a.add("GATE-05-13", "Addon 4 §K", "AI key rotation is specified.", "13", has("docs/13_SECURITY_PRIVACY.md", "key rotation", "rotate"), "key-rotation markers found")

    # GATE-05B — official Addon 5 §M (12)
    addon5_sha = hashlib.sha256(OFFICIAL_ADDON_5.read_bytes()).hexdigest() if OFFICIAL_ADDON_5.is_file() else "missing"
    a5_rows = coverage_rows("A5-", "### 4.6", "## 5.")
    a.add("GATE-05B-01", "Addon 5 §M", "All 14 official Addon 5 rows are integrated with a changelog record.", "00 §4.6 + CHANGELOG + source hash", addon5_sha == EXPECTED_ADDON_5_SHA256 and len(a5_rows) == 14 and all("INTEGRATED" in row for row in a5_rows) and has("docs/CHANGELOG.md", "official Addon 5"), f"source SHA-256={addon5_sha}; rows={len(a5_rows)}")
    a.add("GATE-05B-02", "Addon 5 §M", "Owner handbook includes cadence, report, review, sampling, red flags, stuck choices and oracle.", "30 §§2–10", all(has("docs/30_OWNER_OPERATING_HANDBOOK.md", marker) for marker in ["session cadence", "required session report", "five-minute review", "Phase 0 review", "Spot-check sampling", "Red flags", "stuck", "oracle"]), "all owner-handbook control headings found")
    a.add("GATE-05B-03", "Addon 5 §M", "Evidence matrix/cross-reference and dated evidence convention exist.", "30 + 19 + evidence/", has("docs/30_OWNER_OPERATING_HANDBOOK.md", "evidence/YYYY-MM-DD-<short-task-name>") and has("docs/19_VIBE_CODING_PLAYBOOK.md", "evidence/YYYY-MM-DD-<task>") and (ROOT / "evidence").is_dir(), "evidence convention and directory found")
    a.add("GATE-05B-04", "Addon 5 §M", "Thirteen red flags and response ladder are present.", "30 §8 + 19", has("docs/30_OWNER_OPERATING_HANDBOOK.md", "Response ladder", "| 13 |") and has("docs/19_VIBE_CODING_PLAYBOOK.md", "red flag"), "red-flag ladder markers found")
    a.add("GATE-05B-05", "Addon 5 §M", "Development-time egress policy and S1 incident response are in security contract.", "13 §3.1", has("docs/13_SECURITY_PRIVACY.md", "development-time", "egress", "S1", "cloud agent"), "egress and S1 markers found")
    a.add("GATE-05B-06", "Addon 5 §M", "Golden Month workflow/blessing/regeneration control is defined.", "14 §5.6", has("docs/14_TESTING_QA_PLAN.md", "Golden Month", "BLESSED", "regen-golden"), "golden-month control markers found")
    oracle_ok, oracle_detail = oracle_workbook_check()
    a.add("GATE-05B-07", "Addon 5 §M", "Independent formula-visible oracle workbook and local real-pilot attachment rule exist.", "14 §5.7 + 28 §4.4 + oracle workbook", oracle_ok and has("docs/14_TESTING_QA_PLAN.md", "formula-visible", "independent oracle") and has("docs/28_ACCEPTANCE_UAT_AND_GO_LIVE.md", "isolated local", "never enter the repository"), oracle_detail)
    a.add("GATE-05B-08", "Addon 5 §M", "Hygiene/exclusions, survivability, What's New and no-activation decision are resolved.", ".gitignore + 17 + 24 + 01", has(".gitignore", "sample-data", "client") and has("docs/17_CODING_STANDARDS.md", "git bundle", "generated") and has("docs/24_RELEASE_AND_VERSIONING_RUNBOOK.md", "What's New") and has("docs/01_PRD.md", "no activation", "no expiry"), "repo/release/licensing controls found")
    a.add("GATE-05B-09", "Addon 5 §M", "Stuck protocol includes stop point, rollback and exactly three owner options.", "19 §7.5 + 30 §9", has("docs/19_VIBE_CODING_PLAYBOOK.md", "second failed attempt", "last green", "exactly three") and all(has("docs/30_OWNER_OPERATING_HANDBOOK.md", marker) for marker in ["Simpler approach", "Timeboxed spike", "Descope proposal"]), "stop/rollback/three-option markers found")
    a.add("GATE-05B-10", "Addon 5 §M", "Questionnaire is sendable, tracked and time-limited by an explicit owner decision.", "21 §5.3–5.4 + 28 §6 + 18 OQ-023", has("docs/21_CLIENT_ONBOARDING_QUESTIONNAIRE.md", "Sendable-form", "Response tracker", "OQ-023") and has("docs/28_ACCEPTANCE_UAT_AND_GO_LIVE.md", "owner-approved N-business-day") and has("docs/18_GLOSSARY_ASSUMPTIONS_OPEN_QUESTIONS.md", "OQ-023"), "send/track/deadline-decision markers found")
    a.add("GATE-05B-11", "Addon 5 §M", "Post-go-live incident-to-release/request-to-backlog operations loop exists.", "23 + 24 + 27 + 28", has("docs/23_CONSULTANT_HANDOVER_AND_SUPPORT.md", "S1", "S4", "backlog") and has("docs/24_RELEASE_AND_VERSIONING_RUNBOOK.md", "What's New") and has("docs/27_BACKLOG.md", "post-go-live") and has("docs/28_ACCEPTANCE_UAT_AND_GO_LIVE.md", "Operations loop"), "support/release/backlog-loop markers found")
    a.add("GATE-05B-12", "Addon 5 §M", "Convergence/freeze notice and Addon 5 §N next actions are recorded.", "00 + 19 + 16", has("docs/00_INDEX.md", "frozen contract", "numbered addon") and has("docs/19_VIBE_CODING_PLAYBOOK.md", "Convergence") and has("docs/16_ROADMAP_PHASES.md", "Addon 5 §N", "STOP"), "freeze and next-action markers found")

    # Auxiliary non-gate safeguards: these are reported but do not inflate the 70-check count.
    links_ok, links_detail = local_markdown_links()
    no_code_ok, no_code_detail = no_product_code_check()
    client_egress_ok = (
        has(
            "docs/29_CLIENT_REQUIREMENTS_PACK.md",
            "Please do not email, upload",
            "metadata only",
            "isolated local client/consultant pilot environment",
        )
        and has(
            "docs/21_CLIENT_ONBOARDING_QUESTIONNAIRE.md",
            "No real-file transmission",
            "Do not email, upload or share the files",
            "metadata-only column layout",
        )
        and all(
            forbidden not in text("docs/29_CLIENT_REQUIREMENTS_PACK.md")
            for forbidden in (
                "Provide the prior-year general-ledger extract if it exists",
                "Provide both; the tool maintains them from then on",
            )
        )
    )
    client_egress_detail = (
        "client pack and sendable questionnaire require metadata/style-only discovery and isolated-local pilot use"
        if client_egress_ok
        else "client-facing egress wording is missing its safe boundary or contains a known unsafe request"
    )
    auxiliary = {
        "Internal Markdown links": (links_ok, links_detail),
        "No product code before approval": (no_code_ok, no_code_detail),
        "Official Addon 5 SHA-256": (addon5_sha == EXPECTED_ADDON_5_SHA256, addon5_sha),
        "Client-facing development-time egress copy": (client_egress_ok, client_egress_detail),
    }
    return a.results, auxiliary


def render(results: list[Result], auxiliary: dict[str, tuple[bool, str]]) -> str:
    grouped: dict[str, list[Result]] = {}
    for result in results:
        grouped.setdefault(result.check_id.rsplit("-", 1)[0], []).append(result)
    total_pass = sum(result.passed for result in results)
    lines = [
        "# Phase 0 six-source re-audit — 2026-10-02",
        "",
        "## Scope and limits",
        "",
        "This is an evidence-backed **Phase 0 documentation re-audit** of the 70 checklist rows in",
        "`docs/14_TESTING_QA_PLAN.md` §15: Kickoff (9), Addon 1 §O (12), Addon 2 §I (12),",
        "Addon 3 §J (12), Addon 4 §K (13), and official Addon 5 §M (12). It validates the specified",
        "documents, safe fixtures and recorded documentation walkthrough evidence. It does **not** claim a",
        "product build, installer, Windows run, pilot, UAT, or go-live has occurred; those are explicitly",
        "post-approval / later-phase activities.",
        "",
        f"- Official Addon 5 SHA-256 expected: `{EXPECTED_ADDON_5_SHA256}`",
        f"- Results: **{total_pass}/{len(results)} PASS**",
        "- Phase 0 approval: **not recorded**; this report is evidence for presentation, not approval.",
        "",
        "## Gate results",
        "",
    ]
    for gate, gate_results in grouped.items():
        source = gate_results[0].source
        passed = sum(item.passed for item in gate_results)
        lines.extend([f"### {gate} — {source}: {passed}/{len(gate_results)} PASS", "", "| ID | Assertion | Result | Evidence | Detail |", "|---|---|---|---|---|"])
        for item in gate_results:
            result = "PASS" if item.passed else "FAIL"
            lines.append(f"| `{item.check_id}` | {item.assertion} | **{result}** | {item.evidence} | {item.detail} |")
        lines.append("")
    lines.extend(["## Auxiliary safeguards (not part of the 70)", "", "| Check | Result | Detail |", "|---|---|---|"])
    for label, (passed, detail) in auxiliary.items():
        lines.append(f"| {label} | **{'PASS' if passed else 'FAIL'}** | {detail} |")
    lines.extend(["", "## Conclusion", "", "A green documentary re-audit establishes only that the Phase 0 specification set is ready to be presented.", "It never substitutes for the exact recorded owner approval required in `CHANGELOG.md` and `SESSION_LOG.md`.", ""])
    return "\n".join(lines)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--report", type=Path, help="write the Markdown report to this repository-relative or absolute path")
    args = parser.parse_args()
    results, auxiliary = audit()
    report = render(results, auxiliary)
    if args.report:
        output = args.report if args.report.is_absolute() else ROOT / args.report
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(report, encoding="utf-8")
    print(report)
    return 0 if all(item.passed for item in results) and all(passed for passed, _ in auxiliary.values()) else 1


if __name__ == "__main__":
    raise SystemExit(main())
