"""Integration tests for ENG-10: Run correlation ID threading and export presence."""

import openpyxl
from starlette.testclient import TestClient

from app.api.main import SESSION_TOKEN, app
from app.engine.exports.excel_pack import (
    ExceptionRow,
    MonthEndPackData,
    PackContext,
    export_excel_pack,
)


def test_api_run_rules_accepts_correlation_and_claim_id():
    """Verify POST /api/v1/exceptions/run accepts correlationId and claimId."""
    client = TestClient(app)

    test_corr_id = "run-corr-int-999"
    test_claim_id = "opencode2-20261005T2106Z-4fdf"

    res = client.post(
        "/api/v1/exceptions/run",
        headers={"X-Session-Token": SESSION_TOKEN},
        json={
            "period": "FY26-P09",
            "asOfDate": "2026-11-12",
            "correlationId": test_corr_id,
            "claimId": test_claim_id,
        },
    )
    assert res.status_code == 200
    data = res.json()["data"]
    assert "rulesRun" in data
    assert data["rulesRun"] == 24


def test_exported_pack_carries_run_correlation_id_in_exception_register(tmp_path):
    """Verify that export_excel_pack threads and outputs Correlation ID and Claim ID in Sheet 5."""
    out_xlsx = tmp_path / "month_end_pack.xlsx"
    ctx = PackContext(
        project_name="Acme Corp Test",
        correlation_id="run-corr-stamp-12345",
        claim_id="claim-id-stamp-67890",
    )

    from decimal import Decimal

    exc_row = ExceptionRow(
        exception_id=1,
        rule_id="EXC-001",
        rule_name="Import imbalance",
        rule_version="1.0",
        family="Import Integrity",
        severity="High",
        mark="!",
        subject="Batch 101",
        subject_key="batch:101",
        entity="IN01",
        account_code="1000",
        account_name="Cash",
        cost_centre="CC100",
        vendor="None",
        period="FY26-P09",
        first_seen="2026-10-01",
        amount_at_risk=Decimal("500.00"),
        status="open",
        owner="GL Accountant",
        age_days=1,
        overdue="No",
        sla_due="2026-10-15",
        effective_threshold="0.00",
        flagged_again="No",
        notes_count=0,
        last_note="None",
        evidence_refs="ref1",
        raised_at="2026-10-01 10:00:00",
        last_seen_at="2026-10-01 10:00:00",
        closed_at=None,
        run_id=12,
        correlation_id=ctx.correlation_id,
        claim_id=ctx.claim_id,
    )

    pack_data = MonthEndPackData(context=ctx, exception_rows=[exc_row])
    export_excel_pack(out_xlsx, pack_data)

    assert out_xlsx.exists()
    wb = openpyxl.load_workbook(out_xlsx)
    assert "Exception Register" in wb.sheetnames
    ws = wb["Exception Register"]

    headers = [cell.value for cell in ws[6] if cell.value is not None]
    assert "Correlation ID" in headers
    assert "Claim ID" in headers

    corr_col = headers.index("Correlation ID") + 1
    claim_col = headers.index("Claim ID") + 1

    assert ws.cell(row=7, column=corr_col).value == "run-corr-stamp-12345"
    assert ws.cell(row=7, column=claim_col).value == "claim-id-stamp-67890"
