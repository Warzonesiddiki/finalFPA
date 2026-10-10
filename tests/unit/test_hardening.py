"""Comprehensive unit tests for Excel and CSV Hardening per docs/04_SOURCE_MAPPING_AND_IMPORT_SPEC.md §8 (X1..X26) and §9 (C1..C12)."""

from pathlib import Path

import openpyxl
import pytest

from app.engine.imports.hardening import (
    detect_csv_encoding_and_delimiter,
    detect_hidden_sheets,
    is_total_subtotal_row,
    load_hardened_excel_sheet,
    read_hardened_csv,
    trim_trailing_empty,
)


@pytest.mark.tst_id("TST-SEC-01")
def test_csv_utf8_bom_and_delimiter_detection(tmp_path: Path):
    """Test C1 (UTF-8 BOM stripping) and C5 (delimiter auto-detection)."""
    csv_file = tmp_path / "test_bom.csv"
    # Write UTF-8 with BOM and semicolon delimiter
    content = "\ufeffAccount;Description;Amount\r\n1001;Cash;500.00\r\n1002;Receivables;1200.00\r\n"
    csv_file.write_bytes(content.encode("utf-8"))

    res = detect_csv_encoding_and_delimiter(csv_file)
    assert res.has_bom is True
    assert res.encoding == "utf-8-sig"
    assert res.delimiter == ";"
    assert any(f.slug == "import.encodingDetected" for f in res.findings)
    assert any(f.slug == "import.delimiterDetected" for f in res.findings)

    # Read hardened CSV
    headers, data_rows, findings = read_hardened_csv(csv_file)
    assert headers == ["Account", "Description", "Amount"]
    assert len(data_rows) == 2
    assert data_rows[0] == ["1001", "Cash", "500.00"]


@pytest.mark.tst_id("TST-SEC-02")
def test_csv_delimiters_comma_tab_pipe(tmp_path: Path):
    """Test C5 auto-detection for tab, pipe, comma."""
    # Tab delimited
    tab_file = tmp_path / "test_tab.csv"
    tab_file.write_text("ColA\tColB\tColC\n1\t2\t3\n4\t5\t6\n", encoding="utf-8")
    res_tab = detect_csv_encoding_and_delimiter(tab_file)
    assert res_tab.delimiter == "\t"

    # Pipe delimited
    pipe_file = tmp_path / "test_pipe.csv"
    pipe_file.write_text("ColA|ColB|ColC\n1|2|3\n4|5|6\n", encoding="utf-8")
    res_pipe = detect_csv_encoding_and_delimiter(pipe_file)
    assert res_pipe.delimiter == "|"

    # Comma delimited with trailing delimiter and whitespace (C11, C12)
    comma_file = tmp_path / "test_comma.csv"
    comma_file.write_text(
        " ColA , ColB , ColC ,\n 10 , 20 , 30 ,\n 40 , 50 , 60 ,\n", encoding="utf-8"
    )
    headers, rows, findings = read_hardened_csv(comma_file)
    assert headers == ["ColA", "ColB", "ColC"]
    assert rows[0] == ["10", "20", "30"]
    assert any(f.slug == "import.trailingDelimiter" for f in findings)


@pytest.mark.tst_id("TST-SEC-03")
def test_csv_unsupported_encoding():
    """Test C4: Unknown/other encoding fails gracefully."""
    invalid_bytes = b"\xff\xfe\x00\x00\x00\x00\x12\x34"  # UTF-32 LE pattern
    with pytest.raises(ValueError) as excinfo:
        detect_csv_encoding_and_delimiter(invalid_bytes)
    assert "import.encodingUnsupported" in str(excinfo.value)


def test_hidden_sheet_detection(tmp_path: Path):
    """Test X7: Hidden sheets detection."""
    wb = openpyxl.Workbook()
    ws_main = wb.active
    ws_main.title = "VisibleData"
    ws_main["A1"] = "Header"
    ws_main["A2"] = "Value"

    ws_hidden = wb.create_sheet("HiddenNotes")
    ws_hidden["A1"] = "Confidential"
    ws_hidden.sheet_state = "hidden"

    wb_path = tmp_path / "test_hidden.xlsx"
    wb.save(wb_path)

    visible, hidden = detect_hidden_sheets(wb)
    assert visible == ["VisibleData"]
    assert hidden == ["HiddenNotes"]

    data = load_hardened_excel_sheet(wb_path)
    assert data.sheet_name == "VisibleData"
    assert "HiddenNotes" in data.hidden_sheets
    assert any(
        f.slug == "import.hiddenSheetSkipped" and f.sheet_name == "HiddenNotes"
        for f in data.findings
    )


def test_merged_cells_header_unmerge_and_data_quarantine(tmp_path: Path):
    """Test X2: Merged cells in header are unmerged and propagated; merged data cells are quarantined."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MergedTest"

    # Row 1: Merged header across A1:B1
    ws["A1"] = "Company"
    ws.merge_cells("A1:B1")
    ws["C1"] = "Amount"

    # Row 2: Normal data
    ws["A2"] = "ACME"
    ws["B2"] = "Division 1"
    ws["C2"] = 1000

    # Row 3: Merged data cell across A3:A4
    ws["A3"] = "BETA"
    ws.merge_cells("A3:A4")
    ws["B3"] = "Division 2"
    ws["C3"] = 2000

    # Row 4: Data
    ws["B4"] = "Division 3"
    ws["C4"] = 3000

    wb_path = tmp_path / "test_merged.xlsx"
    wb.save(wb_path)

    data = load_hardened_excel_sheet(wb_path, header_rows=[1])
    # Headers unmerged and propagated
    assert data.headers == ["Company", "Company", "Amount"]

    # Row 2 is valid data
    assert len(data.rows) == 1
    assert data.rows[0] == ["ACME", "Division 1", 1000]

    # Rows 3 and 4 should be quarantined
    assert len(data.quarantined_rows) >= 2
    quarantined_row_indices = {f.row_index for f in data.quarantined_rows}
    assert 3 in quarantined_row_indices
    assert 4 in quarantined_row_indices
    assert any(f.slug == "import.mergedDataCells" for f in data.quarantined_rows)


def test_formula_cached_verification(tmp_path: Path):
    """Test X11: Formula cells with no cached value are quarantined."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "FormulaTest"

    ws["A1"] = "Code"
    ws["B1"] = "Amount"

    ws["A2"] = "V1"
    ws["B2"] = 100

    ws["A3"] = "V2"
    # Uncalculated formula (openpyxl creates it without cached value by default)
    ws["B3"] = "=B2*2"

    wb_path = tmp_path / "test_formula.xlsx"
    wb.save(wb_path)

    # In openpyxl saving without Excel calculation leaves B3 with None in data_only=True
    data = load_hardened_excel_sheet(wb_path, header_rows=[1])

    # Row 2 is clean
    assert len(data.rows) == 1
    assert data.rows[0] == ["V1", 100]

    # Row 3 is quarantined due to missing cached formula value
    assert len(data.quarantined_rows) == 1
    f = data.quarantined_rows[0]
    assert f.slug == "import.formulaNoCachedValue"
    assert f.row_index == 3
    assert f.cell_ref == "B3"


def test_multi_row_headers_concatenation(tmp_path: Path):
    """Test X3: Multi-row header detection and concatenation with ' / '."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MultiHeader"

    # Row 1 & 2 multi-row headers
    ws["A1"] = "Transaction"
    ws["A2"] = "ID"

    ws["B1"] = "Amount"
    ws["B2"] = "Debit"

    ws["C1"] = "Amount"
    ws["C2"] = "Credit"

    # Row 3: Data
    ws["A3"] = "TX100"
    ws["B3"] = 500
    ws["C3"] = 0

    wb_path = tmp_path / "test_multiheader.xlsx"
    wb.save(wb_path)

    data = load_hardened_excel_sheet(wb_path, header_rows=[1, 2])
    assert data.headers == ["Transaction / ID", "Amount / Debit", "Amount / Credit"]
    assert len(data.rows) == 1
    assert data.rows[0] == ["TX100", 500, 0]
    assert any(f.slug == "import.multiRowHeader" for f in data.findings)


def test_trailing_empty_and_blank_rows(tmp_path: Path):
    """Test X5 (trailing blank rows/cols trimmed) and X6 (blank rows inside data ignored)."""
    grid = [
        ["A", "B", None, ""],
        ["1", "2", None, ""],
        ["", None, None, None],
        ["3", "4", "", ""],
        ["", "", "", ""],
        [None, None, None, None],
    ]
    trimmed, trimmed_rows, trimmed_cols = trim_trailing_empty(grid)
    assert trimmed_rows == 2  # bottom 2 rows trimmed
    assert trimmed_cols == 2  # right 2 columns trimmed
    assert len(trimmed) == 4
    assert len(trimmed[0]) == 2


def test_total_subtotal_row_exclusion(tmp_path: Path):
    """Test X4: Embedded Total / Subtotal row detection and exclusion."""
    assert is_total_subtotal_row(["Total", 50000.00]) is True
    assert is_total_subtotal_row(["Subtotal - Operations", 12500]) is True
    assert is_total_subtotal_row(["Grand Total", "₹1,250,000.00"]) is True
    assert is_total_subtotal_row(["Regular Line", 100]) is False
    assert is_total_subtotal_row(["Total without numbers", None, ""]) is False

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TotalsSheet"

    ws["A1"] = "Account"
    ws["B1"] = "Amount"

    ws["A2"] = "1001"
    ws["B2"] = 1000

    ws["A3"] = "1002"
    ws["B3"] = 2000

    ws["A4"] = "Subtotal"
    ws["B4"] = 3000

    ws["A5"] = "Grand Total"
    ws["B5"] = 3000

    wb_path = tmp_path / "test_totals.xlsx"
    wb.save(wb_path)

    data = load_hardened_excel_sheet(wb_path, header_rows=[1])
    assert len(data.rows) == 2
    assert data.rows[0] == ["1001", 1000]
    assert data.rows[1] == ["1002", 2000]
    assert len(data.ignored_total_rows) == 2
    assert any(f.slug == "import.totalRowsIgnored" and f.row_index == 4 for f in data.findings)
    assert any(f.slug == "import.totalRowsIgnored" and f.row_index == 5 for f in data.findings)
