"""Tests for the flag-gated overflow cascade in ``app.engine.pptx_fill.layout``.

Adapted from ``tests/unit/test_layout_adaptive.py`` of deckforge at the pinned commit,
as ``E13``/``R6`` require of every adopted source. The behaviours carried over are
upstream's: overflow detection, a font-reduction step that respects a per-role floor,
narrower-width reflow, and a final escalation when shrinking is not enough. The harness
is re-expressed — upstream measures IR regions through a ``TextMeasurer`` and a layout
solver, while this module measures a real python-pptx shape with the ``ppt_fit`` budget
(``DEC-064``).

One behaviour is deliberately *not* carried over: upstream's ``_split_slide`` rewrites its
IR into continuation slides. Splitting changes the slide count, which docs/12 §2.1 makes a
contract decision rather than a layout one, so this module reports ``split_needed`` and
leaves the slide alone. The upstream test that asserted split content is therefore
re-expressed as ``test_split_request_never_drops_content``.

The cascade is **off by default**, so every test that exercises it sets the flag.
"""

from __future__ import annotations

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from app.engine.pptx_fill import (
    SLIDE_LAYOUTS,
    add_slide_from_layout,
    open_template,
    resolve_shape,
    set_text,
)
from app.engine.pptx_fill.layout import (
    CASCADE_FLAG_ENV,
    FONT_REDUCTION_STEP_PT,
    MAX_ITERATIONS,
    MIN_BODY_FONT_PT,
    MIN_HEADING_FONT_PT,
    REFLOW_WIDTH_RATIO,
    CascadeResult,
    cascade_enabled,
    detect_overflow,
    handle_overflow,
    measure_lines,
    shape_box_inches,
)


@pytest.fixture
def blank_deck():
    return Presentation()


@pytest.fixture
def enabled(monkeypatch):
    monkeypatch.setenv(CASCADE_FLAG_ENV, "1")
    return True


@pytest.fixture
def disabled(monkeypatch):
    monkeypatch.delenv(CASCADE_FLAG_ENV, raising=False)


def _textbox(prs, name, text, width_in, height_in, font_pt, left_in=0.5, top_in=0.5):
    """A plain text box with exactly one run, for precise geometry control."""
    shape = prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(
        Inches(left_in), Inches(top_in), Inches(width_in), Inches(height_in)
    )
    shape.name = name
    run = shape.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(font_pt)
    return shape


# ---------------------------------------------------------------------------
# The flag — the module must be inert unless a doc-12 amendment turns it on
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("value", ["1", "true", "TRUE", "yes", "on", " On "])
def test_cascade_enabled_for_truthy_values(monkeypatch, value):
    monkeypatch.setenv(CASCADE_FLAG_ENV, value)
    assert cascade_enabled() is True


@pytest.mark.parametrize("value", ["0", "false", "no", "", "maybe"])
def test_cascade_disabled_for_other_values(monkeypatch, value):
    monkeypatch.setenv(CASCADE_FLAG_ENV, value)
    assert cascade_enabled() is False


def test_cascade_is_off_when_the_flag_is_absent(disabled):
    """ADR-013: default OFF, so docs/12 §3.4's trim stays the canonical path."""
    assert cascade_enabled() is False


def test_handle_overflow_is_a_no_op_when_disabled(blank_deck, disabled):
    shape = _textbox(blank_deck, "PPT-999_body", "word " * 400, 4.0, 0.5, 12)
    size_before = shape.text_frame.paragraphs[0].runs[0].font.size
    width_before, _ = shape_box_inches(shape)

    result = handle_overflow(shape)

    assert result == CascadeResult("PPT-999_body", None, False, False)
    assert result.overflowed is False
    assert result.split_needed is False
    assert shape.text_frame.paragraphs[0].runs[0].font.size == size_before
    assert shape_box_inches(shape)[0] == width_before


# ---------------------------------------------------------------------------
# Measurement
# ---------------------------------------------------------------------------
def test_shape_box_inches_reports_geometry(blank_deck):
    shape = _textbox(blank_deck, "PPT-999_body", "x", 4.0, 1.5, 12)
    assert shape_box_inches(shape) == pytest.approx((4.0, 1.5))


class _SizelessShape:
    """python-pptx refuses ``width = None``, but a shape from another producer can
    carry no extent; the geometry helpers must treat that as zero, not crash."""

    name = "PPT-999_body"
    width = None
    height = None

    def has_text_frame(self):
        return True

    @property
    def text_frame(self):
        return self

    @property
    def text(self):
        return "word " * 100


def test_shape_box_inches_is_zero_without_a_size():
    assert shape_box_inches(_SizelessShape()) == (0.0, 0.0)


def test_measure_lines_counts_wrapping(blank_deck):
    shape = _textbox(blank_deck, "PPT-999_body", "a" * 400, 4.0, 3.0, 12)
    assert measure_lines(shape, 12.0, 12, 4.0) > 1


def test_measure_lines_is_zero_for_empty_text(blank_deck):
    shape = _textbox(blank_deck, "PPT-999_body", "", 4.0, 3.0, 12)
    assert measure_lines(shape, 12.0, 12, 4.0) == 0


# ---------------------------------------------------------------------------
# Detection
# ---------------------------------------------------------------------------
def test_no_overflow_when_space_is_sufficient(blank_deck, disabled):
    shape = _textbox(blank_deck, "PPT-999_body", "Short line.", 8.0, 2.0, 12)
    assert detect_overflow(shape, 12.0, 12) is False


def test_overflow_detected_when_text_exceeds_the_frame(blank_deck, disabled):
    shape = _textbox(blank_deck, "PPT-999_body", "word " * 500, 2.0, 0.4, 12)
    assert detect_overflow(shape, 12.0, 12) is True


def test_no_overflow_reported_for_a_sizeless_shape(disabled):
    assert detect_overflow(_SizelessShape(), 12.0, 12) is False


# ---------------------------------------------------------------------------
# Step 1 — font reduction and the floors
# ---------------------------------------------------------------------------
def test_font_reduction_never_goes_below_the_body_floor(blank_deck, enabled):
    shape = _textbox(blank_deck, "PPT-999_body", "word " * 900, 1.5, 0.3, 12)
    handle_overflow(shape, max_lines=4)
    size = shape.text_frame.paragraphs[0].runs[0].font.size
    assert size.pt >= MIN_BODY_FONT_PT


def test_heading_floor_is_higher_than_the_body_floor():
    """A smaller heading is unreadable, not merely tight."""
    assert MIN_HEADING_FONT_PT > MIN_BODY_FONT_PT


def test_headings_shrink_less_far_than_body_copy(blank_deck, enabled):
    body = _textbox(blank_deck, "PPT-999_body", "word " * 900, 1.5, 0.3, 12)
    title = _textbox(blank_deck, "PPT-999_title", "word " * 900, 1.5, 0.3, 12)
    handle_overflow(body, max_lines=4)
    handle_overflow(title, max_lines=4)
    body_size = body.text_frame.paragraphs[0].runs[0].font.size.pt
    title_size = title.text_frame.paragraphs[0].runs[0].font.size.pt
    assert body_size < title_size


def test_font_is_reduced_by_a_bounded_step(blank_deck, enabled):
    """Shrinking is iterative and capped, so a single call cannot vanish the text."""
    shape = _textbox(blank_deck, "PPT-999_body", "word " * 40, 3.0, 0.35, 20)
    assert detect_overflow(shape, 20.0, 2) is True
    handle_overflow(shape, max_lines=2)
    size = shape.text_frame.paragraphs[0].runs[0].font.size.pt
    assert 20.0 - (MAX_ITERATIONS * FONT_REDUCTION_STEP_PT) <= size < 20.0


def test_shape_that_already_fits_is_left_alone(blank_deck, enabled):
    shape = _textbox(blank_deck, "PPT-999_body", "Short line.", 8.0, 2.0, 12)
    size_before = shape.text_frame.paragraphs[0].runs[0].font.size
    result = handle_overflow(shape, max_lines=12)
    assert result.overflowed is False
    assert shape.text_frame.paragraphs[0].runs[0].font.size == size_before


# ---------------------------------------------------------------------------
# Steps 2 and 3 — reflow, then an honest split request
# ---------------------------------------------------------------------------
def test_reflow_narrows_the_shape_when_shrinking_is_not_enough(blank_deck, enabled):
    shape = _textbox(blank_deck, "PPT-999_body", "word " * 900, 1.5, 0.3, 12)
    width_before = shape_box_inches(shape)[0]
    result = handle_overflow(shape, max_lines=4)
    if result.split_needed:
        assert shape_box_inches(shape)[0] < width_before
        assert shape_box_inches(shape)[0] == pytest.approx(
            width_before * REFLOW_WIDTH_RATIO, rel=1e-3
        )


def test_split_is_requested_rather_than_silently_resized(blank_deck, enabled):
    """When nothing fits, the cascade asks for a split and says so."""
    shape = _textbox(blank_deck, "PPT-999_body", "word " * 2000, 1.0, 0.2, 12)
    result = handle_overflow(shape, max_lines=1)
    assert result.split_needed is True
    assert result.overflowed is True
    assert result.font_pt >= MIN_BODY_FONT_PT


def test_split_request_never_drops_content(blank_deck, enabled):
    """Upstream's "split preserves all items", re-expressed for a reported split.

    Nothing is split here, so the guarantee is stronger: the shape keeps its whole text.
    """
    text = "word " * 2000
    shape = _textbox(blank_deck, "PPT-999_body", text, 1.0, 0.2, 12)
    handle_overflow(shape, max_lines=1)
    assert shape.text_frame.text == text


def test_cascade_result_names_the_shape_it_touched(blank_deck, enabled):
    shape = _textbox(blank_deck, "PPT-002_narrative", "word " * 500, 1.5, 0.3, 12)
    result = handle_overflow(shape, max_lines=2)
    assert result.shape_name == "PPT-002_narrative"


# ---------------------------------------------------------------------------
# Integration with the real deck
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("flag", [False, True])
def test_real_deck_blocks_never_overflow(monkeypatch, flag):
    """Whatever the flag, the shipped narrative blocks fit their template frames."""
    if flag:
        monkeypatch.setenv(CASCADE_FLAG_ENV, "1")
    else:
        monkeypatch.delenv(CASCADE_FLAG_ENV, raising=False)

    prs = open_template()
    slide = add_slide_from_layout(prs, SLIDE_LAYOUTS["PPT-002"])
    narrative = resolve_shape(slide, "PPT-002_narrative")
    set_text(narrative, "September landed 5.5% above budget, driven by repairs.")
    assert detect_overflow(narrative, 12.0, 6) is False


def test_cascade_is_wired_into_the_deck_builders():
    """The module is not dead code: §3.4 prose blocks run it on every build."""
    import inspect

    from app.engine.exports import ppt_pack

    source = inspect.getsource(ppt_pack._write_prose)
    assert "handle_overflow" in source


def test_deck_build_is_unaffected_by_the_flag():
    """Turning the flag on must not change the bytes of a default deck."""
    import hashlib

    from app.engine.exports.ppt_pack import generate_powerpoint_deck

    def digest():
        prs = generate_powerpoint_deck()
        return hashlib.sha256(
            "".join(slide._element.xml for slide in prs.slides).encode("utf-8")
        ).hexdigest()

    import os

    previous = os.environ.get(CASCADE_FLAG_ENV)
    try:
        os.environ.pop(CASCADE_FLAG_ENV, None)
        without = digest()
        os.environ[CASCADE_FLAG_ENV] = "1"
        with_flag = digest()
    finally:
        os.environ.pop(CASCADE_FLAG_ENV, None)
        if previous is not None:
            os.environ[CASCADE_FLAG_ENV] = previous

    assert without == with_flag