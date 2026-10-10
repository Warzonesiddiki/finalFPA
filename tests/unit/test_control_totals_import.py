"""End-to-end IMP-025 ingestion from an Excel ControlTotals worksheet."""

from decimal import Decimal

from openpyxl import Workbook, load_workbook

from app.engine.imports.parser import (
    parse_excel_transactions,
    prescan_file,
)
from app.engine.imports.profiles import BUILTIN_PROFILES
from app.engine.rules.rules_catalog_001_008 import evaluate_catalog_exc_003
from app.engine.store.db import DatabaseManager
from app.engine.store.exceptions_repo import ExceptionsRepository
from app.engine.store.import_repo import ImportRepository


def _workbook(path, *, include_control_totals=True):
    workbook = Workbook()
    data = workbook.active
    data.title = "Data"
    data.append(
        [
            "Voucher",
            "PostingDate",
            "CompanyCode",
            "MainAccount",
            "CostCenter",
            "Debit",
            "Credit",
            "Currency",
        ]
    )
    data.append(
        [
            "VCH-1",
            "2026-09-22",
            "IN01",
            "5300",
            "CC-110",
            1000,
            650,
            "INR",
        ]
    )
    if include_control_totals:
        totals = workbook.create_sheet("ControlTotals")
        totals.append(["Scope", "Measure", "SuppliedTotal", "Tolerance"])
        totals.append(["gl_control_total", "net", Decimal("0.00"), Decimal("0.00")])
    workbook.save(path)
    workbook.close()
    return path


def _assert_rejected_without_facts(project_dir, batch, transactions):
    database = DatabaseManager(project_dir=project_dir)
    batch_id = ImportRepository(database).commit_batch(batch, transactions)

    sqlite_conn = database.get_sqlite_connection()
    try:
        status = sqlite_conn.execute(
            "SELECT status FROM FactImportBatch WHERE batch_id = ?", [batch_id]
        ).fetchone()[0]
    finally:
        sqlite_conn.close()

    duck_conn = database.get_duckdb_connection()
    try:
        fact_count = duck_conn.execute(
            "SELECT COUNT(*) FROM FactActual WHERE import_batch_id = ?", [batch_id]
        ).fetchone()[0]
    finally:
        duck_conn.close()

    assert status == "rejected"
    assert fact_count == 0


def test_prescan_and_parse_excel_read_data_sheet_and_control_totals(tmp_path):
    path = _workbook(tmp_path / "actuals.xlsx")

    prescan = prescan_file(path)
    batch, transactions = parse_excel_transactions(
        path,
        profile=BUILTIN_PROFILES[0],
        balance_tolerance=Decimal("500.00"),
    )

    assert prescan.sheet_names == ["Data", "ControlTotals"]
    assert prescan.sample_headers[0] == "Voucher"
    assert batch.file_name == "actuals.xlsx"
    assert batch.sheet_name == "Data"
    assert batch.file_checksum == prescan.file_checksum
    assert len(transactions) == 1
    assert transactions[0].source_row_ref == "Data!2"
    assert transactions[0].net_amount == Decimal("350.00")

    check = next(item for item in batch.checks if item.check_code == "IMP-025")
    assert check.status == "fail"
    assert check.sample_rows[0]["loaded_total"] == "350.00"
    assert check.sample_rows[0]["variance"] == "350.00"
    assert check.sample_rows[0]["source_row_ref"] == "ControlTotals!2"


def test_unaccepted_control_total_variance_rejects_without_committing_facts(tmp_path):
    path = _workbook(tmp_path / "unaccepted.xlsx")
    database = DatabaseManager(project_dir=tmp_path / "project")
    batch, transactions = parse_excel_transactions(
        path,
        profile=BUILTIN_PROFILES[0],
        balance_tolerance=Decimal("500.00"),
    )

    batch_id = ImportRepository(database).commit_batch(batch, transactions)
    sqlite_conn = database.get_sqlite_connection()
    try:
        batch_status = sqlite_conn.execute(
            "SELECT status FROM FactImportBatch WHERE batch_id = ?", [batch_id]
        ).fetchone()[0]
    finally:
        sqlite_conn.close()
    duck_conn = database.get_duckdb_connection()
    try:
        fact_count = duck_conn.execute(
            "SELECT COUNT(*) FROM FactActual WHERE import_batch_id = ?", [batch_id]
        ).fetchone()[0]
    finally:
        duck_conn.close()

    assert batch_status == "rejected"
    assert fact_count == 0
    rejected_context = ExceptionsRepository(database).build_rule_context("FY26-P09")
    assert evaluate_catalog_exc_003(rejected_context) == []


def test_accepted_control_total_is_persisted_and_raises_exc_003(tmp_path):
    path = _workbook(tmp_path / "accepted.xlsx")
    database = DatabaseManager(project_dir=tmp_path / "project")
    batch, transactions = parse_excel_transactions(
        path,
        profile=BUILTIN_PROFILES[0],
        balance_tolerance=Decimal("500.00"),
        control_total_acceptance={
            "accepted_by": "controller.test",
            "reason": "Documented legacy source variance",
        },
    )
    check = next(item for item in batch.checks if item.check_code == "IMP-025")
    assert check.status == "warn"
    assert check.sample_rows[0]["accepted"] is True
    assert check.sample_rows[0]["accepted_by"] == "controller.test"
    assert check.sample_rows[0]["acceptance_reason"] == "Documented legacy source variance"

    batch_id = ImportRepository(database).commit_batch(batch, transactions)
    context = ExceptionsRepository(database).build_rule_context("FY26-P09")
    findings = evaluate_catalog_exc_003(context)

    assert context.control_totals[0]["batch_id"] == batch_id
    assert context.control_totals[0]["accepted_by"] == "controller.test"
    assert len(findings) == 1
    assert findings[0].subject_key == f"{batch_id}|gl_control_total"
    assert findings[0].amount_at_risk == Decimal("350.00")


def test_missing_control_totals_sheet_is_recorded_as_skipped(tmp_path):
    path = _workbook(tmp_path / "without-totals.xlsx", include_control_totals=False)

    batch, _transactions = parse_excel_transactions(
        path,
        profile=BUILTIN_PROFILES[0],
    )

    check = next(item for item in batch.checks if item.check_code == "IMP-025")
    assert check.status == "skipped"
    assert check.skip_reason == "no control-totals block supplied"


def test_hidden_control_totals_sheet_blocks_commit(tmp_path):
    path = _workbook(tmp_path / "hidden-totals.xlsx")
    workbook = load_workbook(path)
    workbook["ControlTotals"].sheet_state = "hidden"
    workbook.save(path)
    workbook.close()

    batch, _transactions = parse_excel_transactions(
        path,
        profile=BUILTIN_PROFILES[0],
        balance_tolerance=Decimal("500.00"),
    )

    check = next(item for item in batch.checks if item.check_code == "IMP-025")
    assert check.status == "fail"
    assert "hidden" in check.detail
    assert batch.can_commit is False


def test_excel_workbook_with_header_only_data_is_blocked(tmp_path):
    workbook = Workbook()
    data = workbook.active
    data.title = "Data"
    data.append(["Voucher", "PostingDate", "CompanyCode", "MainAccount", "Debit", "Credit"])
    path = tmp_path / "empty-data.xlsx"
    workbook.save(path)
    workbook.close()

    batch, transactions = parse_excel_transactions(path, profile=BUILTIN_PROFILES[0])

    assert transactions == []
    check = next(item for item in batch.checks if item.check_code == "IMP-008")
    assert check.status == "fail"
    assert batch.can_commit is False
    _assert_rejected_without_facts(tmp_path / "empty-data-project", batch, transactions)


def test_excel_missing_voucher_column_is_reported_as_imp005(tmp_path):
    workbook = Workbook()
    data = workbook.active
    data.title = "Data"
    data.append(["PostingDate", "CompanyCode", "MainAccount", "Debit", "Credit"])
    data.append(["2026-09-22", "IN01", "5300", 100, 100])
    path = tmp_path / "missing-voucher.xlsx"
    workbook.save(path)
    workbook.close()

    batch, _transactions = parse_excel_transactions(path, profile=BUILTIN_PROFILES[0])

    check = next(item for item in batch.checks if item.check_code == "IMP-005")
    assert check.status == "fail"
    assert "Voucher" in check.detail
    assert batch.is_balanced is True
    assert batch.can_commit is False
    _assert_rejected_without_facts(tmp_path / "missing-voucher-project", batch, _transactions)


def test_excel_duplicate_headers_are_blocked(tmp_path):
    workbook = Workbook()
    data = workbook.active
    data.title = "Data"
    data.append(
        [
            "Voucher",
            "Voucher",
            "PostingDate",
            "CompanyCode",
            "MainAccount",
            "CostCenter",
            "Debit",
            "Credit",
            "Currency",
        ]
    )
    data.append(["VCH-1", "VCH-1", "2026-09-22", "IN01", "5300", "CC-110", 100, 100, "INR"])
    path = tmp_path / "duplicate-header.xlsx"
    workbook.save(path)
    workbook.close()

    batch, _transactions = parse_excel_transactions(path, profile=BUILTIN_PROFILES[0])

    check = next(item for item in batch.checks if item.check_code == "IMP-006")
    assert check.status == "fail"
    assert batch.is_balanced is True
    assert batch.can_commit is False

    _assert_rejected_without_facts(tmp_path / "duplicate-project", batch, _transactions)
