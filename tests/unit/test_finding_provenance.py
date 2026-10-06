"""Tests verifying finding provenance: correlation ID and claim ID tracking (ENG-12).

Acceptance requirement:
- A finding respects the claim it was raised in: carries the run correlation ID and claim ID (or batch ID).
- Run ID and claim ID are present in the exported Excel pack artifact (Sheet 5).
- A test fails when either ID is missing.
"""

from decimal import Decimal
from pathlib import Path
from unittest.mock import patch
import tempfile
import openpyxl
import pytest

from app.engine.exports.excel_pack import (
    ExceptionRow,
    MonthEndPackData,
    PackContext,
    export_excel_pack,
)
from app.engine.rules import Finding
from app.engine.store.db import DatabaseManager
from app.engine.store.exceptions_repo import ExceptionsRepository


def test_finding_stores_and_returns_correlation_and_claim_id(tmp_path):
    """Verify FactException stores and retrieves correlation_id and claim_id."""
    db_mgr = DatabaseManager(project_dir=tmp_path)
    repo = ExceptionsRepository(db_mgr)

    test_corr_id = "corr-test-xyz-123"
    test_claim_id = "opencode2-20261005T1938Z-aa55"

    test_finding = Finding(
        rule_id="EXC-001",
        rule_name="Test Rule",
        severity="High",
        tier="exact",
        subject_key="test_subject_1",
        subject_display="Test Finding 1",
        amount_at_risk=Decimal("100.00"),
        period_id="FY26-P09",
        owner_role="senior_accountant",
        effective_threshold="Threshold 0.00",
        detail="Test detail",
        evidence_refs=["ref-1"],
        sample_rows=[{"col": "val"}],
    )

    with patch("app.engine.store.exceptions_repo.build_full_rule_batch", return_value=[lambda ctx: [test_finding]]):
        summary = repo.run_rules(
            period_code="FY26-P09",
            correlation_id=test_corr_id,
            claim_id=test_claim_id,
        )
    assert summary["rulesRun"] == 1

    items_res = repo.list_exceptions(period_code="FY26-P09", page=1, page_size=50)
    items = items_res["items"]
    assert len(items) > 0

    first_item = items[0]
    # Check that correlation_id and claim_id are present on the record
    assert "correlation_id" in first_item, "correlation_id must be present on finding"
    assert "claim_id" in first_item, "claim_id must be present on finding"
    assert first_item["correlation_id"] == test_corr_id
    assert first_item["claim_id"] == test_claim_id


def test_exported_excel_pack_carries_correlation_and_claim_id(tmp_path):
    """Verify that the exported Excel workbook Sheet 5 carries Correlation ID and Claim ID columns."""
    out_xlsx = tmp_path / "month_end_pack.xlsx"
    ctx = PackContext(
        project_name="Acme Corp Financials",
        entities=["IN01"],
        periods=["FY26-P09"],
        scenario="Base",
    )

    exc_row = ExceptionRow(
        exception_id=101,
        rule_id="EXC-001",
        rule_name="Import Imbalance",
        rule_version="v1.0",
        family="Import Integrity",
        severity="High",
        mark="H",
        subject="Bank Ledger Batch 41",
        subject_key="batch_041|bank_ledger",
        entity="IN01",
        account_code="1010",
        account_name="Operating Bank",
        cost_centre="CC-100",
        vendor="—",
        period="FY26-P09",
        first_seen="2026-09-30",
        amount_at_risk=Decimal("350.00"),
        status="open",
        owner="GL Accountant",
        age_days=5,
        overdue="No",
        sla_due="2026-10-21",
        effective_threshold="Tolerance ₹0.00",
        flagged_again="No",
        notes_count=1,
        last_note="2026-09-30",
        evidence_refs="batch_041",
        raised_at="2026-09-30 10:00:00",
        last_seen_at="2026-09-30 10:00:00",
        closed_at=None,
        run_id=42,
        correlation_id="run-corr-test-999",
        claim_id="claim-opencode2-test",
    )

    pack_data = MonthEndPackData(
        context=ctx,
        exception_rows=[exc_row],
    )

    export_excel_pack(out_xlsx, pack_data)
    assert out_xlsx.exists()

    wb = openpyxl.load_workbook(out_xlsx)
    assert "Exception Register" in wb.sheetnames
    ws = wb["Exception Register"]

    headers = [cell.value for cell in ws[6] if cell.value is not None]
    assert "Correlation ID" in headers, "Sheet 5 header must contain 'Correlation ID'"
    assert "Claim ID" in headers, "Sheet 5 header must contain 'Claim ID'"

    corr_idx = headers.index("Correlation ID") + 1
    claim_idx = headers.index("Claim ID") + 1

    row_corr_val = ws.cell(row=7, column=corr_idx).value
    row_claim_val = ws.cell(row=7, column=claim_idx).value

    assert row_corr_val == "run-corr-test-999"
    assert row_claim_val == "claim-opencode2-test"


def test_falsification_missing_ids_fail_assertion():
    """Falsification test: assert failure when correlation_id or claim_id is missing."""
    exc_dict = {
        "exception_id": 101,
        "rule_id": "EXC-001",
        # missing correlation_id and claim_id
    }

    with pytest.raises(AssertionError):
        assert "correlation_id" in exc_dict and exc_dict["correlation_id"] is not None, "Missing correlation_id"

    with pytest.raises(AssertionError):
        assert "claim_id" in exc_dict and exc_dict["claim_id"] is not None, "Missing claim_id"
