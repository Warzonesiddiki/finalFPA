"""Unit tests for Excel Pack Generation per 11_EXCEL_OUTPUT_SPEC.md."""

from pathlib import Path

import openpyxl
import pytest

from app.engine.exports import formats as fmt
from app.engine.exports.excel_pack import (
    TAB_COLORS,
    MonthEndPackData,
    PackContext,
    export_excel_pack,
    generate_month_end_pack,
)
from app.engine.exports.stamps import STAMP_FIELDS


@pytest.mark.tst_id("TST-XL-02")
def test_tab_colors():
    """Verify that tab colors match the specification."""
    wb = generate_month_end_pack()
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        expected_color = TAB_COLORS[sheet_name]
        assert ws.sheet_properties.tabColor is not None
        assert ws.sheet_properties.tabColor.rgb.endswith(expected_color)


@pytest.mark.tst_id("TST-XL-03")
def test_values_only_no_formulas():
    """Verify Rule 3.1: Zero formula cells anywhere in the workbook."""
    wb = generate_month_end_pack()
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows(values_only=False):
            for cell in row:
                if cell.value is not None:
                    str_val = str(cell.value).strip()
                    assert not str_val.startswith("="), (
                        f"Formula found in {sheet_name} at {cell.coordinate}: {cell.value}"
                    )


def test_frozen_panes():
    """Verify frozen panes rules: None on Cover, C7 on data sheets."""
    wb = generate_month_end_pack()
    # Cover
    assert wb["Cover & Context"].freeze_panes is None

    # Data sheets
    data_sheets = [
        "Executive Summary & BvA",
        "P&L Statement Analysis",
        "Transaction Detail Drilldown",
        "Exception Register",
        "Forecast Summary",
        "Import Reconciliation",
    ]
    for name in data_sheets:
        ws = wb[name]
        assert ws.freeze_panes == "C7", f"Freeze pane on {name} is {ws.freeze_panes}, expected C7"


def test_autofilters():
    """Verify autofilter rules: None on Cover, enabled on row 6 on data sheets."""
    wb = generate_month_end_pack()
    assert wb["Cover & Context"].auto_filter.ref is None

    data_sheets = [
        "Executive Summary & BvA",
        "P&L Statement Analysis",
        "Transaction Detail Drilldown",
        "Exception Register",
        "Forecast Summary",
        "Import Reconciliation",
    ]
    for name in data_sheets:
        ws = wb[name]
        assert ws.auto_filter.ref is not None, f"Autofilter missing on {name}"
        assert ws.auto_filter.ref.startswith("A6:"), (
            f"Autofilter on {name} should start at A6, got {ws.auto_filter.ref}"
        )


def test_gridlines():
    """Verify gridlines: off on Cover, on for data sheets."""
    wb = generate_month_end_pack()
    assert wb["Cover & Context"].views.sheetView[0].showGridLines is False

    data_sheets = [
        "Executive Summary & BvA",
        "P&L Statement Analysis",
        "Transaction Detail Drilldown",
        "Exception Register",
        "Forecast Summary",
        "Import Reconciliation",
    ]
    for name in data_sheets:
        ws = wb[name]
        assert ws.views.sheetView[0].showGridLines is True


def test_cover_stamp_fields_and_defined_names():
    """Verify the 30 stamp fields on Cover & Context rows 7-36 and defined names."""
    wb = generate_month_end_pack()
    ws = wb["Cover & Context"]

    # Verify labels in Column A
    for field_info in STAMP_FIELDS:
        row = field_info.row
        cell_label = ws.cell(row=row, column=1).value
        cell_val = ws.cell(row=row, column=2).value

        assert cell_label == field_info.label
        assert cell_val is not None
        assert str(cell_val).strip() != "", f"Empty stamp value for {field_info.label}"

        # Check defined name exists
        assert field_info.defined_name in wb.defined_names, (
            f"Defined name {field_info.defined_name} missing from workbook"
        )


def test_header_block_structure():
    """Verify that every sheet carries the 4-row header block and row 5 spacer."""
    wb = generate_month_end_pack()
    for name in wb.sheetnames:
        ws = wb[name]
        # Row 1 title and units
        r1_val = ws.cell(row=1, column=1).value
        assert r1_val is not None
        # Spacer row 5
        assert ws.row_dimensions[5].height == 6, (
            f"Row 5 height on {name} is {ws.row_dimensions[5].height}, expected 6"
        )


def test_accounting_number_formats():
    """Verify that accounting number formats are applied properly."""
    wb = generate_month_end_pack()

    # Sheet 2: Executive Summary Actual cell
    ws_bva = wb["Executive Summary & BvA"]
    actual_cell = ws_bva.cell(row=7, column=6)  # Col F = Actual
    assert actual_cell.number_format == fmt.MONEY_IN
    assert actual_cell.alignment.horizontal == "right"

    # Var % cell
    var_pct_cell = ws_bva.cell(row=7, column=9)  # Col I = Var %
    assert var_pct_cell.number_format == fmt.PCT_1DP
    assert var_pct_cell.alignment.horizontal == "right"

    # Transaction Detail Posting Date cell
    ws_detail = wb["Transaction Detail Drilldown"]
    date_cell = ws_detail.cell(row=7, column=11)  # Posting date
    assert date_cell.number_format == fmt.DATE_DMY


def test_empty_state_handling():
    """Verify that empty datasets render empty state messages and do not crash."""
    empty_data = MonthEndPackData(context=PackContext())
    wb = generate_month_end_pack(empty_data)

    assert wb["Executive Summary & BvA"].cell(row=7, column=1).value == "No data for this filter."
    assert (
        wb["Transaction Detail Drilldown"].cell(row=7, column=1).value
        == "No transactions for this filter."
    )
    assert (
        wb["Exception Register"].cell(row=7, column=1).value
        == "No open exceptions — nothing requires review."
    )
    assert wb["Forecast Summary"].cell(row=7, column=1).value == "No locked forecast version yet."
    assert wb["Import Reconciliation"].cell(row=7, column=1).value == "No import batches yet."


def test_export_excel_pack_atomic(tmp_path: Path):
    """Verify atomic file export and readability."""
    export_file = tmp_path / "test_pack.xlsx"
    res_path = export_excel_pack(export_file)

    assert res_path.exists()
    assert res_path == export_file
    # Check that temp file does not remain
    assert not export_file.with_suffix(".tmp").exists()

    # Load back with openpyxl to verify integrity
    wb_loaded = openpyxl.load_workbook(export_file)
    assert len(wb_loaded.sheetnames) == 7
    assert "Cover & Context" in wb_loaded.sheetnames
    assert "Pack_Stamp_Version" in wb_loaded.defined_names
